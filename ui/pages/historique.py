"""
ui/pages/historique.py — Historique mensuel + export
"""
import tkinter.filedialog as fd
import customtkinter as ctk
from config import C, MONTHS_FR
from ui.components import make_card, show_toast


class HistoriquePage:
    # Pagination : on charge 24 mois à la fois
    PAGE_SIZE = 24

    def render(self, container: ctk.CTkFrame, app):
        db = app.db
        # État de pagination (combien de mois on a affiché)
        self._displayed = self.PAGE_SIZE
        self._app = app

        ctk.CTkLabel(container, text="📋  Historique",
                     font=ctk.CTkFont(size=24, weight="bold"),
                     text_color=C["text"]).grid(
            row=0, column=0, sticky="w", padx=24, pady=(20, 6)
        )

        scroll = ctk.CTkScrollableFrame(container, fg_color=C["bg"])
        scroll.grid(row=1, column=0, sticky="nsew", padx=24, pady=(0, 16))
        scroll.grid_columnconfigure(0, weight=1)
        container.grid_rowconfigure(1, weight=1)
        self._scroll = scroll

        # Charge la première page
        rows = db.monthly_summary(self._displayed)

        if not rows:
            ctk.CTkLabel(scroll,
                         text="Aucune donnée enregistrée.\nCommencez par 'Saisir le mois' !",
                         text_color=C["muted"],
                         font=ctk.CTkFont(size=14),
                         justify="center").grid(row=0, column=0, pady=80)
            return

        # ── En-tête ─────────────────────────────────────────
        hdr = ctk.CTkFrame(scroll, fg_color=C["sidebar"], corner_radius=8)
        hdr.grid(row=0, column=0, sticky="ew", pady=(0, 4))
        hdr.grid_columnconfigure((0, 1, 2, 3, 4), weight=1)
        for i, h in enumerate(["Mois", "Revenus", "Dépenses", "Épargne", "Bilan"]):
            ctk.CTkLabel(hdr, text=h,
                         font=ctk.CTkFont(size=12, weight="bold"),
                         text_color="white").grid(
                row=0, column=i, padx=16, pady=10, sticky="w"
            )

        # ── Conteneur des lignes (sera repeuplé par "Charger plus") ──
        self._rows_frame = ctk.CTkFrame(scroll, fg_color="transparent")
        self._rows_frame.grid(row=1, column=0, sticky="ew")
        self._rows_frame.grid_columnconfigure(0, weight=1)
        self._render_rows(rows)

        # ── Bouton "Charger plus" + boutons d'export ────────
        self._actions_frame = ctk.CTkFrame(scroll, fg_color="transparent")
        self._actions_frame.grid(row=2, column=0, sticky="ew", pady=14)
        self._actions_frame.grid_columnconfigure(0, weight=1)
        self._render_actions()

        # ── Boutons export (déclarés ici pour rester dans le scope) ─
        ef = self._actions_frame

        def export_csv():
            path = fd.asksaveasfilename(
                defaultextension=".csv",
                filetypes=[("CSV", "*.csv")],
                initialfile="finance_tracker.csv",
            )
            if path:
                db.export_csv(path)
                show_toast(app, "✅  Export CSV réussi !")

        def export_excel():
            try:
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment
            except ImportError:
                import subprocess, sys
                subprocess.run(
                    [sys.executable, "-m", "pip", "install", "openpyxl"],
                    capture_output=True,
                )
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment

            path = fd.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel", "*.xlsx")],
                initialfile="finance_tracker.xlsx",
            )
            if not path:
                return

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Résumé mensuel"
            hfill = PatternFill("solid", fgColor="1A2340")
            hfont = Font(bold=True, color="FFFFFF", size=11)
            headers = ["Année", "Mois", "Revenus (€)", "Dépenses (€)",
                       "Épargne (€)", "Bilan (€)"]
            for ci, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=ci, value=h)
                cell.fill = hfill
                cell.font = hfont
                cell.alignment = Alignment(horizontal="center")

            for ri, r in enumerate(reversed(list(db.monthly_summary(120))), 2):
                bilan = r["rev"] - r["exp"]
                ws.cell(ri, 1, r["year"])
                ws.cell(ri, 2, MONTHS_FR[r["month"] - 1])
                ws.cell(ri, 3, round(r["rev"], 2))
                ws.cell(ri, 4, round(r["exp"], 2))
                ws.cell(ri, 5, round(r["sav"], 2))
                c = ws.cell(ri, 6, round(bilan, 2))
                c.font = Font(color="16A34A" if bilan >= 0 else "EF4444", bold=True)

            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = (
                    max(len(str(cell.value or "")) for cell in col) + 4
                )
            wb.save(path)
            show_toast(app, "✅  Export Excel réussi !")

        ctk.CTkButton(ef, text="📄  Exporter CSV", command=export_csv,
                      height=36, fg_color=C["muted"],
                      hover_color="#475569").pack(side="right", padx=(6, 0))
        ctk.CTkButton(ef, text="📊  Exporter Excel", command=export_excel,
                      height=36, fg_color="#16A34A",
                      hover_color="#15803D").pack(side="right")

    # ─────────────────────────────────────────────────────────
    #  Rendu des lignes (utilisé à l'init et pour "Charger plus")
    # ─────────────────────────────────────────────────────────
    def _render_rows(self, rows):
        """Rend les lignes du tableau dans self._rows_frame (vide d'abord)."""
        for w in self._rows_frame.winfo_children():
            w.destroy()
        for idx, r in enumerate(rows):
            bilan = r["rev"] - r["exp"]
            bg    = C["card"] if idx % 2 == 0 else C["light"]
            rf    = ctk.CTkFrame(self._rows_frame, fg_color=bg, corner_radius=6)
            rf.grid(row=idx, column=0, sticky="ew", pady=1)
            rf.grid_columnconfigure((0, 1, 2, 3, 4), weight=1)

            ctk.CTkLabel(rf,
                         text=f"{MONTHS_FR[r['month']-1]} {r['year']}",
                         font=ctk.CTkFont(size=13, weight="bold"),
                         text_color=C["text"]).grid(
                row=0, column=0, padx=16, pady=10, sticky="w"
            )
            for col, (val, color) in enumerate([
                (f"{r['rev']:,.0f} €",        C["green"]),
                (f"{r['exp']:,.0f} €",        C["red"]),
                (f"{r['sav']:,.0f} €",        C["blue"]),
                (f"{bilan:+,.0f} €", C["green"] if bilan >= 0 else C["red"]),
            ], start=1):
                ctk.CTkLabel(rf, text=val,
                             text_color=color,
                             font=ctk.CTkFont(
                                 size=13,
                                 weight="bold" if col == 4 else "normal"
                             )).grid(row=0, column=col, padx=16, pady=10, sticky="w")
        # Mémorise le nombre de lignes affichées pour la pagination
        self._rendered_count = len(rows)

    # ─────────────────────────────────────────────────────────
    #  Bouton "Charger plus"
    # ─────────────────────────────────────────────────────────
    def _render_actions(self):
        """(Re)peuple la zone d'actions : bouton 'Charger plus' à gauche."""
        # On ne touche PAS aux boutons d'export à droite (déjà packés).
        # On ajoute / met à jour uniquement le bouton "Charger plus".
        if hasattr(self, "_load_more_btn") and self._load_more_btn is not None:
            try:
                self._load_more_btn.destroy()
            except Exception:
                pass
            self._load_more_btn = None

        db = self._app.db
        # Y a-t-il plus de mois à afficher ?
        next_size = self._displayed + self.PAGE_SIZE
        more_rows = db.monthly_summary(next_size)
        has_more  = len(more_rows) > self._displayed

        if has_more:
            self._load_more_btn = ctk.CTkButton(
                self._actions_frame,
                text=f"⬇  Charger {self.PAGE_SIZE} mois de plus",
                height=36, fg_color=C["primary"],
                hover_color=C["primary"],
                command=self._load_more,
            )
            self._load_more_btn.pack(side="left")

    def _load_more(self):
        """Charge la page suivante de l'historique."""
        self._displayed += self.PAGE_SIZE
        rows = self._app.db.monthly_summary(self._displayed)
        self._render_rows(rows)
        self._render_actions()
