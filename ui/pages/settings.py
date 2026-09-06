"""
ui/pages/settings.py — Paramètres complets :
  import Notion, catégories, email SMTP, PDF, export impôts, mode sombre
"""
import os
import threading
import datetime
import tkinter.filedialog as fd

import customtkinter as ctk
import auth as Auth
import cloud_auth
from config import C, DEFAULT_CATEGORIES, DB_PATH
from logger import configure_log_level, LOG_LEVELS, get_log_file_path
from ui.components import collapsible_card, show_toast
from ui.dialogs import RecurringManagerDialog
from secrets_vault import get_secret, save_secret, delete_secret, keyring_available


class SettingsPage:
    def render(self, container: ctk.CTkFrame, app):
        db = app.db

        container.grid_columnconfigure(0, weight=1)
        container.grid_rowconfigure(1, weight=1)

        ctk.CTkLabel(container, text="⚙️  Paramètres",
                     font=ctk.CTkFont(size=24, weight="bold"),
                     text_color=C["text"]).grid(
            row=0, column=0, sticky="w", padx=24, pady=(20, 6)
        )

        scroll = ctk.CTkScrollableFrame(container, fg_color=C["bg"])
        scroll.grid(row=1, column=0, sticky="nsew", padx=24, pady=(0, 16))
        scroll.grid_columnconfigure(0, weight=1)

        row = [0]
        def next_row():
            r = row[0]; row[0] += 1; return r

        # ══════════════════════════════════════════════════════
        #  1. RAPPORTS & EXPORTS
        # ══════════════════════════════════════════════════════
        rc = collapsible_card(
            scroll, next_row(),
            title="📄  Rapports & Exports",
            subtitle="Générez des rapports PDF ou exportez vos données pour la fiscalité.",
            expanded=False, key="rapports", db=db,
        )

        # ── PDF mensuel ──────────────────────────────────────
        pdf_sec = ctk.CTkFrame(rc, fg_color=C["light"], corner_radius=8)
        pdf_sec.pack(fill="x", padx=16, pady=(0, 8))
        pdf_inner = ctk.CTkFrame(pdf_sec, fg_color="transparent")
        pdf_inner.pack(fill="x", padx=14, pady=10)

        ctk.CTkLabel(pdf_inner, text="Rapport PDF mensuel",
                     font=ctk.CTkFont(size=13, weight="bold"),
                     text_color=C["text"]).grid(row=0, column=0, sticky="w")

        # Sélecteurs mois/année pour le PDF
        from config import MONTHS_FR
        today = datetime.date.today()
        pdf_month_var = ctk.StringVar(value=MONTHS_FR[app.sel_month - 1])
        pdf_year_var  = ctk.StringVar(value=str(app.sel_year))
        sel_row = ctk.CTkFrame(pdf_inner, fg_color="transparent")
        sel_row.grid(row=1, column=0, sticky="w", pady=(6, 0))
        ctk.CTkLabel(sel_row, text="Mois :", font=ctk.CTkFont(size=11),
                     text_color=C["muted"]).pack(side="left", padx=(0, 4))
        ctk.CTkOptionMenu(sel_row, values=MONTHS_FR, variable=pdf_month_var,
                          width=130, height=30).pack(side="left", padx=(0, 8))
        ctk.CTkOptionMenu(sel_row, values=[str(y) for y in range(2020, today.year + 2)],
                          variable=pdf_year_var, width=90, height=30).pack(side="left")

        pdf_status_lbl = ctk.CTkLabel(pdf_inner, text="",
                                       font=ctk.CTkFont(size=11), text_color=C["green"])
        pdf_status_lbl.grid(row=2, column=0, sticky="w", pady=(4, 0))

        def gen_pdf():
            try:
                from utils_pdf import generate_monthly_report
            except ImportError:
                pdf_status_lbl.configure(text="❌  Installez reportlab : pip install reportlab")
                return

            m_num = MONTHS_FR.index(pdf_month_var.get()) + 1
            y_num = int(pdf_year_var.get())
            save_path = fd.asksaveasfilename(
                title="Enregistrer le rapport PDF",
                defaultextension=".pdf",
                filetypes=[("PDF", "*.pdf")],
                initialfile=f"Finance_{y_num}_{m_num:02d}.pdf",
            )
            if not save_path:
                return
            pdf_status_lbl.configure(text="⏳  Génération en cours…")

            def _run():
                try:
                    generate_monthly_report(db, y_num, m_num, save_path)
                    app.after(0, lambda: pdf_status_lbl.configure(text="✅  PDF généré avec succès"))
                    app.after(0, lambda: show_toast(app, "Rapport PDF généré !"))
                except Exception as e:
                    app.after(0, lambda msg=str(e): pdf_status_lbl.configure(text=f"❌  Erreur : {msg}"))
            threading.Thread(target=_run, daemon=True).start()

        ctk.CTkButton(pdf_inner, text="📄  Générer le rapport PDF",
                      height=34, width=200,
                      command=gen_pdf,
                      font=ctk.CTkFont(size=12)).grid(row=3, column=0, sticky="w", pady=(8, 0))

        # ── Export fiscal ─────────────────────────────────────
        tax_sep = ctk.CTkFrame(rc, fg_color="#E2E8F0", height=1)
        tax_sep.pack(fill="x", padx=16, pady=4)

        tax_sec = ctk.CTkFrame(rc, fg_color=C["light"], corner_radius=8)
        tax_sec.pack(fill="x", padx=16, pady=(0, 16))
        tax_inner = ctk.CTkFrame(tax_sec, fg_color="transparent")
        tax_inner.pack(fill="x", padx=14, pady=10)

        ctk.CTkLabel(tax_inner, text="Export fiscal (déclaration impôts)",
                     font=ctk.CTkFont(size=13, weight="bold"),
                     text_color=C["text"]).pack(anchor="w")
        ctk.CTkLabel(tax_inner,
                     text="Plus-values, dons déductibles, frais de formation…",
                     font=ctk.CTkFont(size=11), text_color=C["muted"]).pack(anchor="w")

        tax_year_var  = ctk.StringVar(value=str(today.year - 1))

        tax_row = ctk.CTkFrame(tax_inner, fg_color="transparent")
        tax_row.pack(anchor="w", pady=(6, 0))
        ctk.CTkLabel(tax_row, text="Année :", font=ctk.CTkFont(size=11),
                     text_color=C["muted"]).pack(side="left", padx=(0, 4))
        ctk.CTkOptionMenu(tax_row, values=[str(y) for y in range(2020, today.year + 1)],
                          variable=tax_year_var, width=100, height=30).pack(side="left")

        tax_status_lbl = ctk.CTkLabel(tax_inner, text="",
                                       font=ctk.CTkFont(size=11), text_color=C["green"])
        tax_status_lbl.pack(anchor="w", pady=(4, 0))

        def gen_tax():
            try:
                from utils_taxes import export_tax_summary
            except ImportError:
                tax_status_lbl.configure(text="❌  Module utils_taxes introuvable")
                return
            save_path = fd.asksaveasfilename(
                title="Enregistrer l'export fiscal",
                defaultextension=".csv",
                filetypes=[("CSV", "*.csv")],
                initialfile=f"Fiscal_{tax_year_var.get()}.csv",
            )
            if not save_path:
                return
            tax_status_lbl.configure(text="⏳  Export en cours…")

            def _run():
                try:
                    result = export_tax_summary(db, int(tax_year_var.get()), save_path)
                    app.after(0, lambda n=len(result): tax_status_lbl.configure(text=f"✅  Exporté ({n} lignes)"))
                    app.after(0, lambda: show_toast(app, "Export fiscal généré !"))
                except Exception as e:
                    app.after(0, lambda msg=str(e): tax_status_lbl.configure(text=f"❌  {msg}"))
            threading.Thread(target=_run, daemon=True).start()

        ctk.CTkButton(tax_inner, text="📊  Exporter pour les impôts",
                      height=34, width=220,
                      fg_color="#8B5CF6", hover_color="#7C3AED",
                      command=gen_tax,
                      font=ctk.CTkFont(size=12)).pack(anchor="w", pady=(8, 0))

        # ── Toggle saisie Excel ───────────────────────────────
        excel_sep = ctk.CTkFrame(rc, fg_color="#E2E8F0", height=1)
        excel_sep.pack(fill="x", padx=16, pady=4)

        excel_sec = ctk.CTkFrame(rc, fg_color=C["light"], corner_radius=8)
        excel_sec.pack(fill="x", padx=16, pady=(0, 16))
        excel_inner = ctk.CTkFrame(excel_sec, fg_color="transparent")
        excel_inner.pack(fill="x", padx=14, pady=10)

        excel_var = ctk.BooleanVar(value=db.get_setting("excel_import_enabled", "0") == "1")

        excel_hdr = ctk.CTkFrame(excel_inner, fg_color="transparent")
        excel_hdr.pack(fill="x")

        ctk.CTkLabel(excel_hdr,
                     text="📊  Saisie via template Excel",
                     font=ctk.CTkFont(size=13, weight="bold"),
                     text_color=C["text"]).pack(side="left")

        def _toggle_excel():
            db.set_setting("excel_import_enabled", "1" if excel_var.get() else "0")
            show_toast(app, "✅  Réglage Excel sauvegardé")

        ctk.CTkSwitch(excel_hdr, text="",
                      variable=excel_var,
                      command=_toggle_excel,
                      onvalue=True, offvalue=False).pack(side="right")

        ctk.CTkLabel(excel_inner,
                     text=(
                         "Activez pour afficher les boutons « Télécharger le template » et\n"
                         "« Importer » dans l'en-tête du tableau de bord (mois non clôturés).\n"
                         "Workflow : téléchargez le template → remplissez → importez."
                     ),
                     font=ctk.CTkFont(size=11),
                     text_color=C["muted"],
                     justify="left").pack(anchor="w", pady=(4, 0))

        # ══════════════════════════════════════════════════════
        #  2. EMAIL SMTP
        # ══════════════════════════════════════════════════════
        ec = collapsible_card(
            scroll, next_row(),
            title="📧  Envoi d'email (synthèse mensuelle)",
            subtitle="Configurez votre compte SMTP pour recevoir des résumés mensuels par email.",
            expanded=False, key="email_smtp", db=db,
        )

        # Champs SMTP
        smtp_fields = ctk.CTkFrame(ec, fg_color="transparent")
        smtp_fields.pack(fill="x", padx=20, pady=(0, 6))
        smtp_fields.grid_columnconfigure(1, weight=1)

        def _smtp_row(row_i, label, key, placeholder, show="", initial_value=None):
            ctk.CTkLabel(smtp_fields, text=label, font=ctk.CTkFont(size=11),
                         text_color=C["muted"], width=130, anchor="w").grid(
                row=row_i, column=0, sticky="w", pady=4)
            val = initial_value if initial_value is not None else db.get_setting(key, "")
            var = ctk.StringVar(value=val)
            entry = ctk.CTkEntry(smtp_fields, textvariable=var, height=32,
                                  placeholder_text=placeholder, show=show)
            entry.grid(row=row_i, column=1, sticky="ew", padx=(8, 0), pady=4)
            return var

        smtp_host_var  = _smtp_row(0, "Serveur SMTP",    "smtp_host",  "smtp.gmail.com")
        smtp_port_var  = _smtp_row(1, "Port",             "smtp_port",  "587")
        smtp_user_var  = _smtp_row(2, "Email (login)",    "smtp_user",  "votre@email.com")
        smtp_pass_var  = _smtp_row(3, "Mot de passe",     "smtp_pass",  "App password…", "●",
                                   initial_value=get_secret("smtp_pass") or db.get_setting("smtp_pass", ""))
        smtp_to_var    = _smtp_row(4, "Destinataire",     "smtp_to",    "votre@email.com")

        tls_var = ctk.BooleanVar(value=db.get_setting("smtp_tls", "1") == "1")
        ctk.CTkCheckBox(smtp_fields, text="Utiliser TLS (recommandé)",
                        variable=tls_var,
                        font=ctk.CTkFont(size=11)).grid(
            row=5, column=1, sticky="w", padx=(8, 0), pady=6)

        smtp_status_lbl = ctk.CTkLabel(ec, text="",
                                        font=ctk.CTkFont(size=11), text_color=C["green"])
        smtp_status_lbl.pack(anchor="w", padx=20)

        def save_smtp():
            db.set_setting("smtp_host", smtp_host_var.get().strip())
            db.set_setting("smtp_port", smtp_port_var.get().strip())
            db.set_setting("smtp_user", smtp_user_var.get().strip())
            # Mot de passe : trousseau OS en priorité, fallback DB si keyring indisponible
            pwd_raw = smtp_pass_var.get().strip()
            if not save_secret("smtp_pass", pwd_raw):
                db.set_setting("smtp_pass", pwd_raw)  # fallback DB si keyring absent
            else:
                db.set_setting("smtp_pass", "")       # effacer de la DB après migration
            db.set_setting("smtp_to",   smtp_to_var.get().strip())
            db.set_setting("smtp_tls",  "1" if tls_var.get() else "0")
            smtp_status_lbl.configure(text="✅  Paramètres SMTP sauvegardés")
            show_toast(app, "Configuration email sauvegardée")

        def send_now():
            save_smtp()
            try:
                from utils_email import send_monthly_summary
            except ImportError:
                smtp_status_lbl.configure(text="❌  Module utils_email introuvable")
                return

            host = db.get_setting("smtp_host", "")
            port = int(db.get_setting("smtp_port", "587") or "587")
            user = db.get_setting("smtp_user", "")
            pwd  = get_secret("smtp_pass") or db.get_setting("smtp_pass", "")
            to   = db.get_setting("smtp_to", "")
            tls  = db.get_setting("smtp_tls", "1") == "1"

            if not all([host, user, pwd, to]):
                smtp_status_lbl.configure(text="❌  Renseignez tous les champs SMTP")
                return

            smtp_status_lbl.configure(text="⏳  Envoi en cours…")

            def _run():
                ok, msg = send_monthly_summary(
                    db, app.sel_year, app.sel_month,
                    to, host, port, user, pwd, tls
                )
                icon = "✅" if ok else "❌"
                app.after(0, lambda i=icon, m=msg: smtp_status_lbl.configure(text=f"{i}  {m}"))
            threading.Thread(target=_run, daemon=True).start()

        btn_row = ctk.CTkFrame(ec, fg_color="transparent")
        btn_row.pack(anchor="w", padx=20, pady=(6, 16))
        ctk.CTkButton(btn_row, text="💾  Sauvegarder SMTP",
                      height=34, width=180,
                      fg_color=C["primary"],
                      command=save_smtp,
                      font=ctk.CTkFont(size=12)).pack(side="left", padx=(0, 8))
        ctk.CTkButton(btn_row, text="📧  Envoyer un test maintenant",
                      height=34, width=220,
                      fg_color=C["green"],
                      command=send_now,
                      font=ctk.CTkFont(size=12)).pack(side="left")

        # ══════════════════════════════════════════════════════
        #  3. APPARENCE
        # ══════════════════════════════════════════════════════
        ac = collapsible_card(scroll, next_row(), title="🎨  Apparence",
                               expanded=True, key="apparence", db=db)

        dark_row = ctk.CTkFrame(ac, fg_color=C["light"], corner_radius=8)
        dark_row.pack(fill="x", padx=16, pady=(4, 16))
        dark_inner = ctk.CTkFrame(dark_row, fg_color="transparent")
        dark_inner.pack(fill="x", padx=14, pady=12)

        dark_var = ctk.BooleanVar(value=db.get_setting("dark_mode", "0") == "1")

        def on_dark_toggle():
            app.toggle_dark_mode(dark_var.get())

        ctk.CTkLabel(dark_inner,
                     text="🌙  Mode sombre",
                     font=ctk.CTkFont(size=13, weight="bold"),
                     text_color=C["text"]).pack(side="left")
        ctk.CTkSwitch(dark_inner, text="",
                      variable=dark_var,
                      command=on_dark_toggle,
                      onvalue=True, offvalue=False).pack(side="right")

        ctk.CTkLabel(ac,
                     text="Le changement de thème sera appliqué au prochain démarrage.",
                     font=ctk.CTkFont(size=11), text_color=C["muted"]).pack(
            anchor="w", padx=20, pady=(0, 10))

        # ══════════════════════════════════════════════════════
        #  4. JOURNALISATION (LOGS)
        # ══════════════════════════════════════════════════════
        lc = collapsible_card(scroll, next_row(), title="📋  Journalisation",
                               expanded=False, key="journalisation", db=db)

        log_inner = ctk.CTkFrame(lc, fg_color=C["light"], corner_radius=8)
        log_inner.pack(fill="x", padx=16, pady=(4, 6))
        log_inner_f = ctk.CTkFrame(log_inner, fg_color="transparent")
        log_inner_f.pack(fill="x", padx=14, pady=12)

        # Toggle activation
        log_en_var = ctk.BooleanVar(value=db.get_setting("log_enabled", "1") == "1")
        ctk.CTkLabel(log_inner_f, text="Activer les logs",
                     font=ctk.CTkFont(size=13, weight="bold"),
                     text_color=C["text"]).pack(side="left")
        ctk.CTkSwitch(log_inner_f, text="", variable=log_en_var).pack(side="right")

        # Niveau de log
        level_row = ctk.CTkFrame(lc, fg_color="transparent")
        level_row.pack(anchor="w", padx=20, pady=(4, 0))
        ctk.CTkLabel(level_row, text="Niveau :",
                     font=ctk.CTkFont(size=12), text_color=C["muted"]).pack(side="left")
        log_level_var = ctk.StringVar(value=db.get_setting("log_level", "INFO"))
        ctk.CTkOptionMenu(level_row, values=LOG_LEVELS, variable=log_level_var,
                          height=30, width=130,
                          font=ctk.CTkFont(size=12)).pack(side="left", padx=8)

        # Chemin du fichier log
        log_path = get_log_file_path()
        ctk.CTkLabel(lc, text=f"📂  {log_path}",
                     font=ctk.CTkFont(size=10), text_color=C["muted"]).pack(
            anchor="w", padx=20, pady=(4, 4))

        log_status_lbl = ctk.CTkLabel(lc, text="",
                                       font=ctk.CTkFont(size=11), text_color=C["green"])
        log_status_lbl.pack(anchor="w", padx=20)

        def save_log_settings():
            level   = log_level_var.get()
            enabled = log_en_var.get()
            db.set_setting("log_level",   level)
            db.set_setting("log_enabled", "1" if enabled else "0")
            configure_log_level(level, enabled)
            log_status_lbl.configure(text="✅  Paramètres logs sauvegardés")
            show_toast(app, "Configuration logs sauvegardée")

        ctk.CTkButton(lc, text="💾  Sauvegarder", height=32, width=160,
                      font=ctk.CTkFont(size=12),
                      command=save_log_settings).pack(anchor="w", padx=20, pady=(4, 16))

        # ══════════════════════════════════════════════════════
        #  5. SÉCURITÉ
        # ══════════════════════════════════════════════════════
        sc = collapsible_card(scroll, next_row(), title="🔐  Sécurité & Accès",
                               expanded=True, key="securite", db=db)

        is_cloud = db.get_setting("db_mode", "local") == "online"
        if is_cloud:
            cloud_email = db.get_setting("supabase_user_email", "")
            ctk.CTkLabel(sc, text=f"☁️  Compte cloud : {cloud_email}",
                         font=ctk.CTkFont(size=11),
                         text_color=C["muted"]).pack(anchor="w", padx=20, pady=(0, 4))
        else:
            session_lbl = Auth.session_expiry_label(db)
            if session_lbl:
                ctk.CTkLabel(sc, text=f"🔑  {session_lbl}",
                             font=ctk.CTkFont(size=11),
                             text_color=C["muted"]).pack(anchor="w", padx=20, pady=(0, 4))

        # ── Indicateur trousseau OS ────────────────────────────
        if keyring_available():
            vault_txt   = "🔐  Trousseau Windows actif — secrets chiffrés par DPAPI"
            vault_color = C["green"]
        else:
            vault_txt   = "⚠️  Trousseau indisponible — secrets stockés en base (moins sécurisé)"
            vault_color = C.get("amber", "#F59E0B")
        ctk.CTkLabel(sc, text=vault_txt,
                     font=ctk.CTkFont(size=11),
                     text_color=vault_color).pack(anchor="w", padx=20, pady=(0, 10))

        # ── Changer le mot de passe ───────────────────────────
        pwd_sec = ctk.CTkFrame(sc, fg_color=C["light"], corner_radius=8)
        pwd_sec.pack(fill="x", padx=16, pady=(0, 8))
        pwd_inner = ctk.CTkFrame(pwd_sec, fg_color="transparent")
        pwd_inner.pack(fill="x", padx=14, pady=12)
        pwd_inner.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(pwd_inner, text="Modifier le mot de passe",
                     font=ctk.CTkFont(size=13, weight="bold"),
                     text_color=C["text"]).grid(row=0, column=0, sticky="w")

        show_pwd_form = {"open": False, "widgets": []}

        def _toggle_pwd_form():
            if show_pwd_form["open"]:
                for w in show_pwd_form["widgets"]:
                    try: w.destroy()
                    except Exception: pass
                show_pwd_form["widgets"].clear()
                show_pwd_form["open"] = False
                toggle_pwd_btn.configure(text="✏️  Changer le mot de passe")
                return
            show_pwd_form["open"] = True
            toggle_pwd_btn.configure(text="✕  Annuler")

            lbl1 = ctk.CTkLabel(pwd_inner, text="Nouveau mot de passe",
                                font=ctk.CTkFont(size=11, weight="bold"),
                                text_color=C["muted"])
            lbl1.grid(row=2, column=0, sticky="w", pady=(10, 2))
            npwd_var = ctk.StringVar()
            npwd_entry = ctk.CTkEntry(pwd_inner, textvariable=npwd_var,
                                      show="●", height=36,
                                      font=ctk.CTkFont(size=12),
                                      fg_color=C["card"],
                                      border_color=C["border"])
            npwd_entry.grid(row=3, column=0, sticky="ew")
            npwd_entry.focus()

            lbl2 = ctk.CTkLabel(pwd_inner, text="Confirmer",
                                font=ctk.CTkFont(size=11, weight="bold"),
                                text_color=C["muted"])
            lbl2.grid(row=4, column=0, sticky="w", pady=(8, 2))
            cpwd_var = ctk.StringVar()
            cpwd_entry = ctk.CTkEntry(pwd_inner, textvariable=cpwd_var,
                                      show="●", height=36,
                                      font=ctk.CTkFont(size=12),
                                      fg_color=C["card"],
                                      border_color=C["border"])
            cpwd_entry.grid(row=5, column=0, sticky="ew")

            err_lbl = ctk.CTkLabel(pwd_inner, text="",
                                   font=ctk.CTkFont(size=11),
                                   text_color=C["red"])
            err_lbl.grid(row=6, column=0, sticky="w", pady=(4, 0))

            def _save_pwd():
                np_ = npwd_var.get()
                cp_ = cpwd_var.get()
                min_len = 6 if is_cloud else 4
                if len(np_) < min_len:
                    err_lbl.configure(text=f"⚠  Au moins {min_len} caractères.")
                    return
                if np_ != cp_:
                    err_lbl.configure(text="❌  Les mots de passe ne correspondent pas.")
                    return
                if is_cloud:
                    ok, msg = cloud_auth.set_new_password(app.cloud_client, np_)
                    if not ok:
                        err_lbl.configure(text=f"❌  {msg}")
                        return
                else:
                    Auth.change_password(db, np_)
                err_lbl.configure(text="")
                show_toast(app, "✅  Mot de passe modifié !")
                _toggle_pwd_form()

            save_btn = ctk.CTkButton(pwd_inner,
                                     text="✅  Enregistrer",
                                     height=36,
                                     font=ctk.CTkFont(size=12, weight="bold"),
                                     fg_color=C["primary"],
                                     hover_color="#2955C9",
                                     command=_save_pwd)
            save_btn.grid(row=7, column=0, sticky="ew", pady=(10, 0))
            show_pwd_form["widgets"] = [lbl1, npwd_entry, lbl2, cpwd_entry,
                                        err_lbl, save_btn]

        toggle_pwd_btn = ctk.CTkButton(pwd_inner,
                                       text="✏️  Changer le mot de passe",
                                       height=34, width=210,
                                       font=ctk.CTkFont(size=12),
                                       fg_color=C["primary"],
                                       hover_color="#2955C9",
                                       command=_toggle_pwd_form)
        toggle_pwd_btn.grid(row=1, column=0, sticky="w", pady=(8, 0))

        # ── Se déconnecter ────────────────────────────────────
        def _logout():
            if is_cloud:
                cloud_auth.sign_out(app.cloud_client)
                cloud_auth.clear_persisted_session(db)
                # db_mode reste "online" : au prochain démarrage, l'app
                # affiche l'écran de connexion cloud (pas le mode local).
            else:
                Auth.clear_session(db)
            show_toast(app, "Session fermée — connexion requise au prochain démarrage")

        ctk.CTkButton(sc,
                      text="🚪  Fermer la session (verrouiller)",
                      height=34, width=260,
                      font=ctk.CTkFont(size=12),
                      fg_color="#FEE2E2",
                      text_color=C["red"],
                      hover_color="#FECACA",
                      command=_logout).pack(anchor="w", padx=16, pady=(0, 16))

        # ══════════════════════════════════════════════════════
        #  6. IMPORT NOTION
        # ══════════════════════════════════════════════════════
        ic = collapsible_card(
            scroll, next_row(),
            title="📥  Importer depuis Notion (ZIP)",
            subtitle="Sélectionnez les exports ZIP depuis Notion (revenus ou dépenses).\n"
                     "Les données seront ajoutées sans doublons.",
            expanded=False, key="import_notion", db=db,
        )

        self._import_status_lbl = ctk.CTkLabel(ic, text="",
                                               font=ctk.CTkFont(size=12), text_color=C["green"])
        self._import_status_lbl.pack(anchor="w", padx=20)

        btn_row2 = ctk.CTkFrame(ic, fg_color="transparent")
        btn_row2.pack(anchor="w", padx=20, pady=(6, 16))

        def do_import():
            paths = fd.askopenfilenames(
                title="Sélectionner les fichiers ZIP Notion",
                filetypes=[("ZIP", "*.zip"), ("Tous", "*.*")],
            )
            if not paths:
                return
            self._import_status_lbl.configure(text="⏳  Import en cours…")

            def _run():
                total_rev = total_exp = total_sav = 0
                errors = []
                for p in paths:
                    try:
                        r, e, s = db.import_notion_zip(p)
                        total_rev += r; total_exp += e; total_sav += s
                    except Exception as exc:
                        errors.append(str(exc))
                if errors:
                    msg = "⚠️  Erreur : " + " | ".join(errors)
                else:
                    msg = f"✅  {total_rev} revenus, {total_exp} dépenses, {total_sav} épargnes importés"
                app.after(0, lambda m=msg: self._import_status_lbl.configure(text=m))
            threading.Thread(target=_run, daemon=True).start()

        ctk.CTkButton(btn_row2, text="📂  Choisir ZIP(s) et importer",
                      height=38, command=do_import,
                      font=ctk.CTkFont(size=13)).pack(side="left")

        # ══════════════════════════════════════════════════════
        #  7. CATÉGORIES
        # ══════════════════════════════════════════════════════
        cc = collapsible_card(scroll, next_row(), title="🗂️  Catégories de dépenses",
                               expanded=False, key="categories", db=db)

        list_frame = ctk.CTkFrame(cc, fg_color="transparent")
        list_frame.pack(fill="x", padx=20, pady=(0, 8))

        def refresh_cat_list():
            for w in list_frame.winfo_children():
                w.destroy()
            cats = db.get_categories()
            for i, cat in enumerate(cats):
                row_f = ctk.CTkFrame(list_frame,
                                     fg_color=C["card"] if i % 2 == 0 else C["light"],
                                     corner_radius=6)
                row_f.pack(fill="x", pady=1)
                ctk.CTkLabel(row_f, text=cat["name"],
                             font=ctk.CTkFont(size=12),
                             text_color=C["text"]).pack(side="left", padx=14, pady=7)

                def make_del(c=cat):
                    def _do():
                        from ui.components import confirm_delete
                        confirm_delete(app,
                                       lambda: (db.delete_category(c["id"]),
                                                refresh_cat_list(),
                                                show_toast(app, "Catégorie supprimée")),
                                       label=f"« {c['name']} »")
                    return _do
                ctk.CTkButton(row_f, text="✕", width=28, height=26,
                              fg_color="#FEE2E2", text_color=C["red"],
                              hover_color="#FECACA",
                              command=make_del()).pack(side="right", padx=8, pady=4)

        refresh_cat_list()

        add_row = ctk.CTkFrame(cc, fg_color="transparent")
        add_row.pack(fill="x", padx=20, pady=(4, 4))
        new_cat_var = ctk.StringVar()
        ctk.CTkEntry(add_row, textvariable=new_cat_var,
                     placeholder_text="Nouvelle catégorie…",
                     height=34, width=220).pack(side="left", padx=(0, 8))

        def add_cat():
            name = new_cat_var.get().strip()
            if not name:
                return
            db.add_category(name)
            new_cat_var.set("")
            refresh_cat_list()
            show_toast(app, f"Catégorie « {name.upper()} » ajoutée")

        ctk.CTkButton(add_row, text="＋  Ajouter", height=34, width=120,
                      command=add_cat).pack(side="left")

        def reset_cats():
            db.reset_categories()
            refresh_cat_list()
            show_toast(app, "Catégories réinitialisées")

        ctk.CTkButton(cc,
                      text="↺  Réinitialiser les catégories par défaut",
                      height=32, fg_color=C["muted"], hover_color="#475569",
                      font=ctk.CTkFont(size=12),
                      command=reset_cats).pack(anchor="e", padx=20, pady=(4, 14))

        # ══════════════════════════════════════════════════════
        #  8. BASE DE DONNÉES
        # ══════════════════════════════════════════════════════
        dbc = collapsible_card(scroll, next_row(), title="🗄️  Base de données",
                                expanded=False, key="base_donnees", db=db)
        ctk.CTkLabel(dbc,
                     text=f"Chemin : {db.db_path}",
                     text_color=C["muted"], font=ctk.CTkFont(size=11),
                     justify="left").pack(anchor="w", padx=20)

        def export_csv():
            save_path = fd.asksaveasfilename(
                title="Exporter les données CSV",
                defaultextension=".csv",
                filetypes=[("CSV", "*.csv")],
                initialfile="finance_export.csv",
            )
            if save_path:
                db.export_csv(save_path)
                show_toast(app, "Export CSV terminé !")

        def export_excel():
            try:
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment
            except ImportError:
                show_toast(app, "Export Excel indisponible : le module 'openpyxl' n'est pas installé (pip install openpyxl).")
                return
            from config import MONTHS_FR
            save_path = fd.asksaveasfilename(
                title="Exporter les données Excel",
                defaultextension=".xlsx",
                filetypes=[("Excel", "*.xlsx")],
                initialfile="finance_export.xlsx",
            )
            if not save_path:
                return
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Résumé mensuel"
            hfill = PatternFill("solid", fgColor="1A2340")
            hfont = Font(bold=True, color="FFFFFF", size=11)
            for ci, h in enumerate(["Année", "Mois", "Revenus (€)", "Dépenses (€)",
                                     "Épargne (€)", "Bilan (€)"], 1):
                c = ws.cell(row=1, column=ci, value=h)
                c.fill = hfill; c.font = hfont
                c.alignment = Alignment(horizontal="center")
            for ri, r in enumerate(reversed(list(db.monthly_summary(120))), 2):
                bilan = r["rev"] - r["exp"]
                ws.cell(ri, 1, r["year"])
                ws.cell(ri, 2, MONTHS_FR[r["month"] - 1])
                ws.cell(ri, 3, round(r["rev"], 2))
                ws.cell(ri, 4, round(r["exp"], 2))
                ws.cell(ri, 5, round(r["sav"], 2))
                c = ws.cell(ri, 6, round(bilan, 2))
                c.font = Font(color="16A34A" if bilan >= 0 else "EF4444", bold=True)
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = (
                    max(len(str(cell.value or "")) for cell in col) + 4)
            wb.save(save_path)
            show_toast(app, "✅  Export Excel terminé !")

        export_btn_row = ctk.CTkFrame(dbc, fg_color="transparent")
        export_btn_row.pack(anchor="w", padx=20, pady=(8, 16))
        ctk.CTkButton(export_btn_row, text="📤  Exporter CSV complet",
                      height=32, font=ctk.CTkFont(size=12),
                      fg_color=C["primary"],
                      command=export_csv).pack(side="left", padx=(0, 8))
        ctk.CTkButton(export_btn_row, text="📊  Exporter Excel",
                      height=32, font=ctk.CTkFont(size=12),
                      fg_color="#16A34A", hover_color="#15803D",
                      command=export_excel).pack(side="left")

        # ══════════════════════════════════════════════════════
        #  9. TRANSACTIONS RÉCURRENTES
        # ══════════════════════════════════════════════════════
        rrc = collapsible_card(
            scroll, next_row(),
            title="📅  Transactions récurrentes",
            subtitle="Gérez vos loyers, abonnements, salaires… définissez-les une fois,\n"
                     "et appliquez-les chaque mois depuis le tableau de bord.",
            expanded=False, key="recurrentes", db=db,
        )

        rec_count = len(db.get_recurring_transactions())
        ctk.CTkLabel(rrc, text=f"🔄  {rec_count} transaction(s) configurée(s)",
                     font=ctk.CTkFont(size=12), text_color=C["muted"]).pack(
            anchor="w", padx=20, pady=(0, 8))

        ctk.CTkButton(
            rrc, text="⚙  Gérer les récurrentes",
            height=36, width=200,
            font=ctk.CTkFont(size=13),
            command=lambda: RecurringManagerDialog(app, db),
        ).pack(anchor="w", padx=20, pady=(0, 16))

        # ══════════════════════════════════════════════════════
        #  10. SYNCHRONISATION SUPABASE
        # ══════════════════════════════════════════════════════
        sc2 = collapsible_card(scroll, next_row(), title="☁️  Synchronisation Supabase",
                                expanded=False, key="sync_supabase", db=db)

        current_mode = db.get_setting("db_mode", "local")
        mode_lbl_text = "✅  Mode actuel : En ligne (Supabase)" if current_mode == "online" \
                        else "💾  Mode actuel : Local uniquement"
        mode_color    = C["green"] if current_mode == "online" else C["muted"]
        ctk.CTkLabel(sc2, text=mode_lbl_text,
                     font=ctk.CTkFont(size=12),
                     text_color=mode_color).pack(anchor="w", padx=20, pady=(0, 4))
        ctk.CTkLabel(sc2,
                     text="L'activation du mode en ligne se fait uniquement depuis "
                          "l'écran de connexion («Compte cloud») — jamais ici, "
                          "pour ne jamais activer la synchronisation sans avoir "
                          "d'abord récupéré vos données.",
                     font=ctk.CTkFont(size=10), text_color=C["muted"],
                     wraplength=520, justify="left").pack(anchor="w", padx=20, pady=(0, 12))

        fields_frame = ctk.CTkFrame(sc2, fg_color="transparent")
        fields_frame.pack(fill="x", padx=20, pady=(0, 6))
        fields_frame.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(fields_frame, text="URL du projet Supabase",
                     font=ctk.CTkFont(size=12, weight="bold"),
                     text_color=C["muted"]).grid(row=0, column=0, sticky="w", pady=(0, 4))
        url_var2 = ctk.StringVar(value=db.get_setting("supabase_url", ""))
        ctk.CTkEntry(fields_frame, textvariable=url_var2, height=36,
                     font=ctk.CTkFont(size=12),
                     fg_color=C["light"], border_color=C["border"],
                     placeholder_text="https://xxxx.supabase.co").grid(
            row=1, column=0, sticky="ew", pady=(0, 10))

        ctk.CTkLabel(fields_frame, text="Clé anon (publique)",
                     font=ctk.CTkFont(size=12, weight="bold"),
                     text_color=C["muted"]).grid(row=2, column=0, sticky="w", pady=(0, 4))
        key_var2 = ctk.StringVar(value=db.get_setting("supabase_anon_key", ""))
        ctk.CTkEntry(fields_frame, textvariable=key_var2, height=36,
                     font=ctk.CTkFont(size=12),
                     fg_color=C["light"], border_color=C["border"],
                     placeholder_text="eyJ…").grid(row=3, column=0, sticky="ew")

        sync_status_lbl = ctk.CTkLabel(sc2, text="",
                                        font=ctk.CTkFont(size=11),
                                        wraplength=500,
                                        justify="left")
        sync_status_lbl.pack(anchor="w", padx=20, pady=(6, 0))

        def _set_status(msg: str):
            sync_status_lbl.configure(text=msg)
            sync_status_lbl.update()

        btn_row_sync = ctk.CTkFrame(sc2, fg_color="transparent")
        btn_row_sync.pack(anchor="w", padx=20, pady=(10, 16))

        def _test_connection():
            url = url_var2.get().strip()
            key = key_var2.get().strip()
            if not url or not key:
                _set_status("⚠️  Renseignez l'URL et la clé anon.")
                return
            _set_status("⏳  Test en cours…")

            def _run():
                try:
                    client = cloud_auth.build_client(url, key)
                    client.table("categories").select("id").limit(1).execute()

                    def _done():
                        db.set_setting("supabase_url", url)
                        db.set_setting("supabase_anon_key", key)
                        _set_status(
                            "✅  Connexion OK. Déconnectez-vous puis utilisez "
                            "«Compte cloud» à l'écran de connexion pour "
                            "activer le mode en ligne."
                        )
                    app.after(0, _done)
                except Exception as e:
                    app.after(0, lambda msg=str(e): _set_status(f"❌  {msg}"))

            threading.Thread(target=_run, daemon=True).start()

        def _do_force_push(s):
            _set_status("⏳  Upload en cours…")

            def _run():
                msg = s._do_push(force=True)
                app.after(0, lambda: _set_status(msg))

            threading.Thread(target=_run, daemon=True).start()

        def _confirm_and_push(s, safe: bool, warn_msg: str, remote: dict, local: dict):
            import tkinter.messagebox as mb
            title = "Confirmer l'envoi" if safe else "⚠️  Écart important détecté"
            prefix = f"{warn_msg}\n\n" if not safe else ""
            detail = (
                f"{prefix}Distant : {sum(remote.values())} lignes — "
                f"Local : {sum(local.values())} lignes.\n\n"
                f"Ceci va remplacer les données distantes par les données locales. "
                f"Continuer ?"
            )
            if mb.askyesno(title, detail):
                _do_force_push(s)
            else:
                _set_status("Envoi annulé.")

        def _sync_now():
            if not (is_cloud and app.cloud_client and app.cloud_user_id):
                _set_status("⚠️  Connectez-vous d'abord en mode cloud (écran de connexion).")
                return
            _set_status("⏳  Vérification des comptages…")

            def _run():
                try:
                    from sync_supabase import SupabaseSync
                    s = SupabaseSync.from_session(app.cloud_client, db.db_path, app.cloud_user_id)
                    safe, warn_msg, remote, local = s.check_push_safety()
                    app.after(0, lambda: _confirm_and_push(s, safe, warn_msg, remote, local))
                except Exception as e:
                    app.after(0, lambda msg=str(e): _set_status(f"❌  {msg}"))

            threading.Thread(target=_run, daemon=True).start()

        def _switch_local():
            db.set_setting("db_mode", "local")
            show_toast(app, "💾  Mode local activé")
            _set_status("💾  Mode local activé. Redémarrez l'app pour appliquer.")

        ctk.CTkButton(btn_row_sync, text="🔌  Tester la connexion",
                      height=34, font=ctk.CTkFont(size=12),
                      fg_color=C["primary"],
                      command=_test_connection).pack(side="left", padx=(0, 8))
        ctk.CTkButton(btn_row_sync, text="☁️  Sync maintenant",
                      height=34, font=ctk.CTkFont(size=12),
                      fg_color=C["green"],
                      command=_sync_now).pack(side="left", padx=(0, 8))
        ctk.CTkButton(btn_row_sync, text="💾  Passer en local",
                      height=34, font=ctk.CTkFont(size=12),
                      fg_color=C["muted"], hover_color="#475569",
                      command=_switch_local).pack(side="left")

        # ══════════════════════════════════════════════════════
        #  11. INTELLIGENCE ARTIFICIELLE
        # ══════════════════════════════════════════════════════
        aic = collapsible_card(scroll, next_row(), title="🤖  Intelligence Artificielle",
                                expanded=False, key="ia", db=db)
        ctk.CTkLabel(aic,
                     text="Optionnel — enrichit les recommandations avec une analyse IA personnalisée.\n"
                          "Gemini 2.0 Flash Lite (gratuit, sans CB) · Anthropic Haiku · OpenAI GPT-4o-mini.\n"
                          "Gemini : clé API gratuite sur aistudio.google.com → python -m pip install google-genai",
                     text_color=C["muted"], font=ctk.CTkFont(size=12),
                     justify="left", wraplength=560).pack(anchor="w", padx=20, pady=(0, 12))

        ai_fields = ctk.CTkFrame(aic, fg_color="transparent")
        ai_fields.pack(fill="x", padx=20, pady=(0, 6))
        ai_fields.grid_columnconfigure(0, weight=1)

        # Fournisseur
        ctk.CTkLabel(ai_fields, text="Fournisseur",
                     font=ctk.CTkFont(size=12, weight="bold"),
                     text_color=C["muted"]).grid(row=0, column=0, sticky="w", pady=(0, 4))
        ai_provider_var = ctk.StringVar(
            value=db.get_setting("ai_provider", "gemini"))
        ctk.CTkOptionMenu(ai_fields,
                          values=["gemini", "anthropic", "openai"],
                          variable=ai_provider_var,
                          height=34, font=ctk.CTkFont(size=12),
                          fg_color=C["light"], button_color=C["primary"],
                          text_color=C["text"]).grid(row=1, column=0, sticky="w", pady=(0, 10))

        # Clé API — référence directe à l'entry (évite le bug textvariable+show)
        ctk.CTkLabel(ai_fields, text="Clé API",
                     font=ctk.CTkFont(size=12, weight="bold"),
                     text_color=C["muted"]).grid(row=2, column=0, sticky="w", pady=(0, 4))
        ai_key_entry = ctk.CTkEntry(ai_fields, height=36,
                                    show="●", font=ctk.CTkFont(size=12),
                                    fg_color=C["light"], border_color=C["border"],
                                    placeholder_text="AIza… (Gemini) · sk-ant-… (Anthropic) · sk-… (OpenAI)")
        ai_key_entry.grid(row=3, column=0, sticky="ew")
        # Pré-remplir si déjà enregistrée (trousseau OS en priorité, fallback DB)
        _saved_key = get_secret("ai_api_key") or db.get_setting("ai_api_key", "")
        if _saved_key:
            ai_key_entry.insert(0, _saved_key)

        ai_status_lbl = ctk.CTkLabel(aic, text="",
                                     font=ctk.CTkFont(size=11),
                                     wraplength=500, justify="left")
        ai_status_lbl.pack(anchor="w", padx=20, pady=(6, 0))

        def _ai_set_status(msg: str, color: str = C["muted"]):
            ai_status_lbl.configure(text=msg, text_color=color)

        ai_btn_row = ctk.CTkFrame(aic, fg_color="transparent")
        ai_btn_row.pack(anchor="w", padx=20, pady=(10, 16))

        def _ai_save_and_test():
            provider = ai_provider_var.get().strip()
            key      = ai_key_entry.get().strip()
            if not key:
                _ai_set_status("⚠️  Entrez votre clé API.", C["amber"])
                return
            _ai_set_status("⏳  Test en cours…", C["muted"])

            def _run():
                try:
                    from utils_ai import get_ai_recommendations, build_financial_summary
                    test_summary = "Revenus : 3000 €, Dépenses : 2000 €, Épargne : 500 €."
                    ok, text = get_ai_recommendations(test_summary, 1, provider, key)
                    def _done():
                        if ok:
                            db.set_setting("ai_provider", provider)
                            if not save_secret("ai_api_key", key):
                                db.set_setting("ai_api_key", key)  # fallback DB
                            else:
                                db.set_setting("ai_api_key", "")   # effacer de la DB
                            _ai_set_status(
                                f"✅  Connexion {provider.title()} OK — clé enregistrée.",
                                C["green"])
                            show_toast(app, "🤖  IA configurée avec succès")
                        else:
                            _ai_set_status(f"❌  {text}", C["red"])
                    # Ne jamais interroger Tkinter depuis ce thread — `app._closing`
                    # est un simple booléen Python, sûr à lire hors du thread principal.
                    if not getattr(app, "_closing", False):
                        app.after(0, _done)
                except Exception as e:
                    try:
                        if not getattr(app, "_closing", False):
                            app.after(0, lambda msg=str(e): _ai_set_status(f"❌  {msg}", C["red"]))
                    except Exception:
                        pass

            threading.Thread(target=_run, daemon=True).start()

        def _ai_clear():
            db.set_setting("ai_provider", "")
            delete_secret("ai_api_key")
            db.set_setting("ai_api_key",  "")
            ai_key_entry.delete(0, "end")
            _ai_set_status("🗑️  Clé IA supprimée.", C["muted"])

        ctk.CTkButton(ai_btn_row, text="🔌  Tester & Enregistrer",
                      height=34, font=ctk.CTkFont(size=12),
                      fg_color="#7C3AED", hover_color="#6D28D9",
                      command=_ai_save_and_test).pack(side="left", padx=(0, 8))
        ctk.CTkButton(ai_btn_row, text="🗑️  Effacer la clé",
                      height=34, font=ctk.CTkFont(size=12),
                      fg_color=C["muted"], hover_color="#475569",
                      command=_ai_clear).pack(side="left")

        # ── Reset disclaimer IA ────────────────────────────────
        ctk.CTkFrame(aic, fg_color=C["border"], height=1).pack(
            fill="x", padx=20, pady=(8, 8))

        disclaimer_row = ctk.CTkFrame(aic, fg_color="transparent")
        disclaimer_row.pack(fill="x", padx=20, pady=(0, 4))
        ctk.CTkLabel(disclaimer_row,
                     text="🔔  Affichage du disclaimer (données envoyées à Google) :",
                     font=ctk.CTkFont(size=11), text_color=C["muted"]).pack(side="left")
        disclaimer_accepted = db.get_setting("ai_disclaimer_accepted", "0") == "1"
        disclaimer_status = ctk.CTkLabel(disclaimer_row,
                                         text="Accepté ✅" if disclaimer_accepted else "Non accepté ⬜",
                                         font=ctk.CTkFont(size=11),
                                         text_color=C["green"] if disclaimer_accepted else C["muted"])
        disclaimer_status.pack(side="left", padx=(8, 0))

        def _reset_disclaimer():
            db.set_setting("ai_disclaimer_accepted", "0")
            disclaimer_status.configure(text="Non accepté ⬜", text_color=C["muted"])
            show_toast(app, "Disclaimer IA réinitialisé — sera affiché à la prochaine analyse")

        ctk.CTkButton(aic, text="↺  Réinitialiser le disclaimer IA",
                      height=30, width=240,
                      font=ctk.CTkFont(size=11),
                      fg_color=C["muted"], hover_color="#475569",
                      command=_reset_disclaimer).pack(anchor="w", padx=20, pady=(0, 8))

        # ── Mon profil (contexte injecté dans le prompt IA) ───
        ctk.CTkFrame(aic, fg_color=C["border"], height=1).pack(
            fill="x", padx=20, pady=(0, 12))

        ctk.CTkLabel(aic, text="👤  Mon profil",
                     font=ctk.CTkFont(size=13, weight="bold"),
                     text_color=C["text"]).pack(anchor="w", padx=20)
        ctk.CTkLabel(aic,
                     text="Ce texte est ajouté à chaque analyse IA pour personnaliser les conseils.",
                     font=ctk.CTkFont(size=11), text_color=C["muted"]).pack(
            anchor="w", padx=20, pady=(2, 6))

        ctx_var = ctk.StringVar(value=db.get_setting("user_context", ""))
        ctx_entry = ctk.CTkEntry(aic, textvariable=ctx_var, height=36,
                                  font=ctk.CTkFont(size=12),
                                  placeholder_text="Ex : Jeune couple, un seul salaire, à Sartrouville (78)…")
        ctx_entry.pack(fill="x", padx=20, pady=(0, 6))

        ctx_status = ctk.CTkLabel(aic, text="", font=ctk.CTkFont(size=11),
                                   text_color=C["green"])
        ctx_status.pack(anchor="w", padx=20)

        def _save_ctx():
            db.set_setting("user_context", ctx_var.get().strip())
            ctx_status.configure(text="✅  Profil sauvegardé")
            aic.after(2000, lambda: ctx_status.configure(text=""))

        ctk.CTkButton(aic, text="💾  Sauvegarder le profil",
                      height=32, width=200, font=ctk.CTkFont(size=12),
                      fg_color="#7C3AED", hover_color="#6D28D9",
                      command=_save_ctx).pack(anchor="w", padx=20, pady=(4, 16))
