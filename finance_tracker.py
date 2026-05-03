#!/usr/bin/env python3
"""
Fintrack - Application de suivi financier personnel
Développé avec CustomTkinter + SQLite + Matplotlib
"""

import customtkinter as ctk
import sqlite3
import matplotlib
matplotlib.use("TkAgg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from datetime import datetime
import os, csv, zipfile, io, re
import tkinter.filedialog as fd
import tkinter.messagebox as mb

# ─────────────────────────────────────────────
#  CONFIG
# ─────────────────────────────────────────────
APP_TITLE   = "Fintrack"
APP_VERSION = "1.1"
DB_PATH     = os.path.join(os.path.expanduser("~"), "Documents", "finance_tracker.db")

MONTHS_FR = [
    "Janvier", "Février", "Mars", "Avril", "Mai", "Juin",
    "Juillet", "Août", "Septembre", "Octobre", "Novembre", "Décembre"
]

# Mapping pour parser les noms de mois Notion (FR + variantes)
MONTH_NAME_TO_NUM = {
    "janvier": 1, "fevrier": 2, "février": 2, "mars": 3, "avril": 4,
    "mai": 5, "juin": 6, "juillet": 7, "aout": 8, "août": 8,
    "septembre": 9, "octobre": 10, "novembre": 11,
    "decembre": 12, "décembre": 12,
}

ASSET_TYPES = [
    ("Bourse / ETF / PEA", "bourse"),
    ("Immobilier",         "immobilier"),
    ("Crypto",             "crypto"),
    ("Or & Autres",        "or_autres"),
    ("Compte / Épargne",   "compte"),
]
ASSET_LABEL = {t[1]: t[0] for t in ASSET_TYPES}

DEFAULT_CATEGORIES = [
    "Abonnement", "Alimentation", "Autres", "Banque", "Bricolage",
    "Cadeaux", "Dons", "Famille", "Formation", "Impôt",
    "Logement", "Loisir et sorties", "Mutuelle", "PEA",
    "Restaurants", "Santé", "Shopping", "Sport et fitness",
    "Transport", "Voiture", "Voyage", "Vêtements",
]

PALETTE = ["#3B6FE8","#22C55E","#F59E0B","#EF4444","#8B5CF6",
           "#06B6D4","#EC4899","#14B8A6","#F97316","#6366F1",
           "#84CC16","#F43F5E","#0EA5E9","#A78BFA","#FB923C"]

C = {
    "bg":      "#F0F4F8", "sidebar":  "#1A2340", "sidebar2": "#243050",
    "card":    "#FFFFFF", "border":   "#E2E8F0",  "primary":  "#3B6FE8",
    "green":   "#22C55E", "red":      "#EF4444",  "blue":     "#3B82F6",
    "amber":   "#F59E0B", "text":     "#1E293B",  "muted":    "#64748B",
    "light":   "#F8FAFC",
}


# ─────────────────────────────────────────────
#  HELPERS
# ─────────────────────────────────────────────
def parse_notion_month(raw: str):
    """'Septembre 25 (https://...)' → (2025, 9)"""
    clean = raw.split("(")[0].strip()
    parts = clean.split()
    if len(parts) >= 2:
        m = MONTH_NAME_TO_NUM.get(parts[0].lower())
        try:
            y = int(parts[1])
            y = 2000 + y if y < 100 else y
        except ValueError:
            return None, None
        return (y, m) if m else (None, None)
    return None, None

def parse_amount(raw: str) -> float:
    """'3 924,50 €' → 3924.50"""
    cleaned = raw.replace("\u202f", "").replace("\xa0", "").replace("€", "").strip()
    cleaned = cleaned.replace(" ", "").replace(",", ".")
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


# ─────────────────────────────────────────────
#  BASE DE DONNÉES
# ─────────────────────────────────────────────
class Database:
    def __init__(self, path: str):
        os.makedirs(os.path.dirname(path), exist_ok=True)
        self.con = sqlite3.connect(path, check_same_thread=False)
        self.con.row_factory = sqlite3.Row
        self._init_schema()
        self._migrate()
        self._seed_categories()

    def _init_schema(self):
        self.con.executescript("""
            CREATE TABLE IF NOT EXISTS categories (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE,
                is_default INTEGER DEFAULT 0
            );
            CREATE TABLE IF NOT EXISTS months (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                year INTEGER NOT NULL,
                month INTEGER NOT NULL,
                UNIQUE(year, month)
            );
            CREATE TABLE IF NOT EXISTS expenses (
                id          INTEGER PRIMARY KEY AUTOINCREMENT,
                month_id    INTEGER NOT NULL REFERENCES months(id),
                category_id INTEGER NOT NULL REFERENCES categories(id),
                amount      REAL    NOT NULL,
                label       TEXT    DEFAULT '',
                payee       TEXT    DEFAULT ''
            );
            CREATE TABLE IF NOT EXISTS revenues (
                id       INTEGER PRIMARY KEY AUTOINCREMENT,
                month_id INTEGER NOT NULL REFERENCES months(id),
                source   TEXT    NOT NULL,
                amount   REAL    NOT NULL,
                label    TEXT    DEFAULT ''
            );
            CREATE TABLE IF NOT EXISTS savings (
                id       INTEGER PRIMARY KEY AUTOINCREMENT,
                month_id INTEGER NOT NULL REFERENCES months(id),
                account  TEXT    DEFAULT '',
                amount   REAL    NOT NULL,
                label    TEXT    DEFAULT ''
            );
            CREATE TABLE IF NOT EXISTS assets (
                id         INTEGER PRIMARY KEY AUTOINCREMENT,
                year       INTEGER NOT NULL,
                month      INTEGER NOT NULL,
                asset_type TEXT    NOT NULL,
                asset_name TEXT    NOT NULL,
                value      REAL    NOT NULL,
                notes      TEXT    DEFAULT '',
                UNIQUE(year, month, asset_type, asset_name)
            );
        """)
        self.con.commit()

    def _migrate(self):
        """Migrations non-destructives (colonnes ajoutées)."""
        for stmt in [
            "ALTER TABLE expenses ADD COLUMN payee TEXT DEFAULT ''",
        ]:
            try:
                self.con.execute(stmt)
                self.con.commit()
            except Exception:
                pass  # colonne déjà présente

    def _seed_categories(self):
        cur = self.con.cursor()
        if cur.execute("SELECT COUNT(*) FROM categories").fetchone()[0] == 0:
            for name in DEFAULT_CATEGORIES:
                cur.execute("INSERT OR IGNORE INTO categories(name,is_default) VALUES(?,1)", (name,))
            self.con.commit()

    # ── helpers ──
    def _month_id(self, y, m, create=True):
        cur = self.con.cursor()
        row = cur.execute("SELECT id FROM months WHERE year=? AND month=?", (y, m)).fetchone()
        if row:
            return row["id"]
        if create:
            cur.execute("INSERT INTO months(year,month) VALUES(?,?)", (y, m))
            self.con.commit()
            return cur.lastrowid
        return None

    # ── categories ──
    def get_categories(self):
        return self.con.execute("SELECT * FROM categories ORDER BY name").fetchall()

    def add_category(self, name: str):
        self.con.execute("INSERT OR IGNORE INTO categories(name,is_default) VALUES(?,0)", (name,))
        self.con.commit()

    def del_category(self, cid: int):
        self.con.execute("DELETE FROM categories WHERE id=? AND is_default=0", (cid,))
        self.con.commit()

    # ── expenses ──
    def add_expense(self, y, m, cat_id, amount, label="", payee=""):
        mid = self._month_id(y, m)
        self.con.execute(
            "INSERT INTO expenses(month_id,category_id,amount,label,payee) VALUES(?,?,?,?,?)",
            (mid, cat_id, amount, label, payee))
        self.con.commit()

    def del_expense(self, eid):
        self.con.execute("DELETE FROM expenses WHERE id=?", (eid,))
        self.con.commit()

    def get_expenses(self, y, m, cat_filter=None, payee_filter=None):
        mid = self._month_id(y, m, create=False)
        if not mid:
            return []
        q = """SELECT e.id, c.name AS cat, e.amount, e.label, e.payee
               FROM expenses e JOIN categories c ON e.category_id=c.id
               WHERE e.month_id=?"""
        params = [mid]
        if cat_filter and cat_filter not in ("Toutes", "Toutes catégories"):
            q += " AND c.name=?"
            params.append(cat_filter)
        if payee_filter and payee_filter not in ("Tous", "Toutes enseignes"):
            q += " AND e.payee=?"
            params.append(payee_filter)
        q += " ORDER BY c.name, e.payee"
        return self.con.execute(q, params).fetchall()

    def get_expenses_by_cat(self, y, m, payee_filter=None):
        mid = self._month_id(y, m, create=False)
        if not mid:
            return []
        q = """SELECT c.name, SUM(e.amount) as total
               FROM expenses e JOIN categories c ON e.category_id=c.id
               WHERE e.month_id=?"""
        params = [mid]
        if payee_filter and payee_filter not in ("Tous", "Toutes enseignes"):
            q += " AND e.payee=?"
            params.append(payee_filter)
        q += " GROUP BY c.id ORDER BY total DESC"
        return self.con.execute(q, params).fetchall()

    def get_expenses_by_payee(self, y, m, cat_filter=None):
        mid = self._month_id(y, m, create=False)
        if not mid:
            return []
        q = """SELECT e.payee, SUM(e.amount) as total
               FROM expenses e JOIN categories c ON e.category_id=c.id
               WHERE e.month_id=? AND e.payee != ''"""
        params = [mid]
        if cat_filter and cat_filter not in ("Toutes", "Toutes catégories"):
            q += " AND c.name=?"
            params.append(cat_filter)
        q += " GROUP BY e.payee ORDER BY total DESC LIMIT 15"
        return self.con.execute(q, params).fetchall()

    def get_all_payees(self, y=None, m=None):
        if y and m:
            mid = self._month_id(y, m, create=False)
            if not mid:
                return []
            rows = self.con.execute(
                "SELECT DISTINCT payee FROM expenses WHERE month_id=? AND payee!='' ORDER BY payee",
                (mid,)).fetchall()
        else:
            rows = self.con.execute(
                "SELECT DISTINCT payee FROM expenses WHERE payee!='' ORDER BY payee").fetchall()
        return [r["payee"] for r in rows]

    # ── revenues ──
    def add_revenue(self, y, m, source, amount, label=""):
        mid = self._month_id(y, m)
        self.con.execute("INSERT INTO revenues(month_id,source,amount,label) VALUES(?,?,?,?)",
                         (mid, source, amount, label))
        self.con.commit()

    def del_revenue(self, rid):
        self.con.execute("DELETE FROM revenues WHERE id=?", (rid,))
        self.con.commit()

    def get_revenues(self, y, m):
        mid = self._month_id(y, m, create=False)
        if not mid:
            return []
        return self.con.execute(
            "SELECT * FROM revenues WHERE month_id=? ORDER BY source", (mid,)).fetchall()

    # ── savings ──
    def add_saving(self, y, m, account, amount, label=""):
        mid = self._month_id(y, m)
        self.con.execute("INSERT INTO savings(month_id,account,amount,label) VALUES(?,?,?,?)",
                         (mid, account, amount, label))
        self.con.commit()

    def del_saving(self, sid):
        self.con.execute("DELETE FROM savings WHERE id=?", (sid,))
        self.con.commit()

    def get_savings(self, y, m):
        mid = self._month_id(y, m, create=False)
        if not mid:
            return []
        return self.con.execute(
            "SELECT * FROM savings WHERE month_id=? ORDER BY account", (mid,)).fetchall()

    # ── assets ──
    def upsert_asset(self, y, m, atype, name, value, notes=""):
        self.con.execute("""
            INSERT INTO assets(year,month,asset_type,asset_name,value,notes) VALUES(?,?,?,?,?,?)
            ON CONFLICT(year,month,asset_type,asset_name) DO UPDATE
            SET value=excluded.value, notes=excluded.notes
        """, (y, m, atype, name, value, notes))
        self.con.commit()

    def del_asset(self, aid):
        self.con.execute("DELETE FROM assets WHERE id=?", (aid,))
        self.con.commit()

    def get_assets(self, y, m):
        return self.con.execute(
            "SELECT * FROM assets WHERE year=? AND month=? ORDER BY asset_type, asset_name",
            (y, m)).fetchall()

    def get_asset_history(self):
        return self.con.execute(
            "SELECT year, month, SUM(value) as total FROM assets GROUP BY year, month ORDER BY year, month"
        ).fetchall()

    def get_assets_by_type(self, y, m):
        return self.con.execute("""
            SELECT asset_type, SUM(value) as total
            FROM assets WHERE year=? AND month=? GROUP BY asset_type ORDER BY total DESC
        """, (y, m)).fetchall()

    # ── summary ──
    def monthly_summary(self, limit=18):
        return self.con.execute("""
            SELECT m.year, m.month,
                COALESCE((SELECT SUM(amount) FROM revenues WHERE month_id=m.id),0) AS rev,
                COALESCE((SELECT SUM(amount) FROM expenses WHERE month_id=m.id),0) AS exp,
                COALESCE((SELECT SUM(amount) FROM savings  WHERE month_id=m.id),0) AS sav
            FROM months m ORDER BY m.year DESC, m.month DESC LIMIT ?
        """, (limit,)).fetchall()

    # ── export ──
    def export_csv(self, path):
        rows = self.monthly_summary(120)
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(["Année", "Mois", "Revenus (€)", "Dépenses (€)", "Épargne (€)", "Bilan (€)"])
            for r in reversed(rows):
                w.writerow([r["year"], MONTHS_FR[r["month"]-1],
                             f"{r['rev']:.2f}", f"{r['exp']:.2f}",
                             f"{r['sav']:.2f}", f"{r['rev']-r['exp']:.2f}"])

    # ── import Notion ──
    def import_notion_zip(self, zip_path: str) -> tuple[int, int]:
        """
        Importe un ZIP Notion (revenus ou dépenses).
        Retourne (nb_revenus_importés, nb_dépenses_importées).
        """
        rev_count = 0
        exp_count = 0

        def open_inner(outer_path):
            with zipfile.ZipFile(outer_path) as oz:
                for name in oz.namelist():
                    if name.lower().endswith(".zip"):
                        inner_bytes = oz.read(name)
                        with zipfile.ZipFile(io.BytesIO(inner_bytes)) as iz:
                            for iname in iz.namelist():
                                if iname.lower().endswith("_all.csv"):
                                    yield iname, iz.read(iname).decode("utf-8-sig")
                    elif name.lower().endswith("_all.csv"):
                        yield name, oz.read(name).decode("utf-8-sig")

        for fname, content in open_inner(zip_path):
            reader = csv.DictReader(content.splitlines())
            fieldnames = [k.strip() for k in (reader.fieldnames or [])]
            is_income  = "Category" not in fieldnames and "à qui" not in fieldnames
            is_expense = "Category" in fieldnames or "à qui" in fieldnames

            for raw_row in reader:
                row = {k.strip(): (v or "").strip() for k, v in raw_row.items()}
                mois_raw = row.get("Mois", "")
                if not mois_raw:
                    continue
                y, m = parse_notion_month(mois_raw)
                if not y or not m:
                    continue

                amount_raw = row.get("Montant", "")
                amount = parse_amount(amount_raw)
                if amount <= 0:
                    continue

                label = row.get("Nom", row.get("Notes", ""))

                if is_income:
                    source = row.get("Nom", "Revenu")
                    self.add_revenue(y, m, source, amount, "")
                    rev_count += 1

                elif is_expense:
                    cat_name = row.get("Category", "Autres").strip() or "Autres"
                    payee    = row.get("à qui", "").strip()

                    # Crée la catégorie si elle n'existe pas
                    self.add_category(cat_name)
                    cats = {c["name"]: c["id"] for c in self.get_categories()}
                    cat_id = cats.get(cat_name) or cats.get("Autres")
                    if not cat_id:
                        continue

                    self.add_expense(y, m, cat_id, amount, label, payee)
                    exp_count += 1

        return rev_count, exp_count


# ─────────────────────────────────────────────
#  COMPOSANTS UI RÉUTILISABLES
# ─────────────────────────────────────────────
def make_card(parent, **kw):
    d = dict(fg_color=C["card"], corner_radius=14, border_width=1, border_color=C["border"])
    d.update(kw)
    return ctk.CTkFrame(parent, **d)

def kpi_card(parent, title, value, color, icon=""):
    card = make_card(parent)
    inner = ctk.CTkFrame(card, fg_color="transparent")
    inner.pack(fill="both", expand=True, padx=18, pady=14)
    ctk.CTkLabel(inner, text=f"{icon}  {title}", font=ctk.CTkFont(size=11),
                 text_color=C["muted"]).pack(anchor="w")
    ctk.CTkLabel(inner, text=value, font=ctk.CTkFont(size=24, weight="bold"),
                 text_color=color).pack(anchor="w", pady=(2, 0))
    return card

def nav_btn(parent, text, cmd):
    return ctk.CTkButton(parent, text=text, anchor="w", height=42,
                         font=ctk.CTkFont(size=13), fg_color="transparent",
                         hover_color=C["sidebar2"], text_color="#CBD5E1",
                         corner_radius=8, command=cmd)

def add_row_widget(parent, row_idx, cols, delete_cmd):
    bg = C["light"] if row_idx % 2 == 0 else C["card"]
    f = ctk.CTkFrame(parent, fg_color=bg, corner_radius=6)
    f.grid(row=row_idx, column=0, columnspan=len(cols)+1, sticky="ew", pady=1, padx=2)
    for i, (w, text, color) in enumerate(cols):
        f.grid_columnconfigure(i, weight=w)
        ctk.CTkLabel(f, text=text, text_color=color or C["text"],
                     font=ctk.CTkFont(size=12)).grid(row=0, column=i, padx=10, pady=7, sticky="w")
    ctk.CTkButton(f, text="✕", width=28, height=26, fg_color="#FEE2E2",
                  text_color="#EF4444", hover_color="#FECACA",
                  command=delete_cmd).grid(row=0, column=len(cols), padx=8, pady=4, sticky="e")
    return f

def filter_bar(parent, label, values, var, cmd):
    ctk.CTkLabel(parent, text=label, font=ctk.CTkFont(size=11), text_color=C["muted"]).pack(side="left", padx=(0,4))
    ctk.CTkOptionMenu(parent, values=values, variable=var, command=cmd,
                      width=160, height=30, font=ctk.CTkFont(size=12)).pack(side="left", padx=(0,14))


# ─────────────────────────────────────────────
#  APPLICATION PRINCIPALE
# ─────────────────────────────────────────────
class App(ctk.CTk):
    def __init__(self):
        super().__init__()
        ctk.set_appearance_mode("light")
        ctk.set_default_color_theme("blue")
        self.title(f"💰 {APP_TITLE}  v{APP_VERSION}")
        self.geometry("1350x860")
        self.minsize(1050, 680)

        self.db = Database(DB_PATH)

        now = datetime.now()
        self.sel_year  = now.year
        self.sel_month = now.month

        # Filtres dashboard persistants
        self.dash_cat_filter   = "Toutes catégories"
        self.dash_payee_filter = "Toutes enseignes"

        self._build_layout()
        self._go("dashboard")

    # ── layout ──
    def _build_layout(self):
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)

        sb = ctk.CTkFrame(self, width=225, corner_radius=0, fg_color=C["sidebar"])
        sb.grid(row=0, column=0, sticky="nsew")
        sb.grid_propagate(False)
        sb.grid_rowconfigure(8, weight=1)
        sb.grid_columnconfigure(0, weight=1)

        lf = ctk.CTkFrame(sb, fg_color="transparent")
        lf.grid(row=0, column=0, padx=20, pady=(28,18), sticky="ew")
        ctk.CTkLabel(lf, text="💰 Finance", font=ctk.CTkFont(size=21, weight="bold"),
                     text_color="#F1F5F9").pack(anchor="w")
        ctk.CTkLabel(lf, text="Tracker", font=ctk.CTkFont(size=21, weight="bold"),
                     text_color="#60A5FA").pack(anchor="w")

        ctk.CTkFrame(sb, height=1, fg_color="#2D3F5C").grid(
            row=1, column=0, sticky="ew", padx=16, pady=(0,8))

        nav_items = [
            ("dashboard",  "🏠   Tableau de bord"),
            ("monthly",    "📝   Saisir le mois"),
            ("analyses",   "📊   Analyses"),
            ("patrimoine", "📈   Patrimoine"),
            ("historique", "📋   Historique"),
            ("settings",   "⚙️   Paramètres"),
        ]
        self._nav_btns = {}
        for i, (page, label) in enumerate(nav_items, start=2):
            btn = nav_btn(sb, label, lambda p=page: self._go(p))
            btn.grid(row=i, column=0, padx=12, pady=2, sticky="ew")
            self._nav_btns[page] = btn

        ctk.CTkLabel(sb, text=f"v{APP_VERSION}", font=ctk.CTkFont(size=10),
                     text_color="#3D5278").grid(row=9, column=0, padx=20, pady=14, sticky="sw")

        self._content = ctk.CTkFrame(self, corner_radius=0, fg_color=C["bg"])
        self._content.grid(row=0, column=1, sticky="nsew")
        self._content.grid_columnconfigure(0, weight=1)
        self._content.grid_rowconfigure(0, weight=1)

    def _go(self, page):
        for name, btn in self._nav_btns.items():
            btn.configure(fg_color=C["primary"] if name==page else "transparent",
                          text_color="#FFFFFF" if name==page else "#CBD5E1")
        for w in self._content.winfo_children():
            w.destroy()
        plt.close("all")
        getattr(self, f"_page_{page}")()

    # ─── PAGE: DASHBOARD ───
    def _page_dashboard(self):
        y, m = self.sel_year, self.sel_month

        # Filtres courants
        cat_f   = self.dash_cat_filter
        payee_f = self.dash_payee_filter

        scroll = ctk.CTkScrollableFrame(self._content, fg_color="transparent")
        scroll.grid(row=0, column=0, sticky="nsew", padx=24, pady=20)
        scroll.grid_columnconfigure((0,1,2,3), weight=1)

        # ── Titre + sélecteur de mois ──
        hdr = ctk.CTkFrame(scroll, fg_color="transparent")
        hdr.grid(row=0, column=0, columnspan=4, sticky="ew", pady=(0,12))
        hdr.grid_columnconfigure(1, weight=1)

        left = ctk.CTkFrame(hdr, fg_color="transparent")
        left.grid(row=0, column=0, sticky="w")
        ctk.CTkLabel(left, text="Tableau de bord",
                     font=ctk.CTkFont(size=24, weight="bold"), text_color=C["text"]).pack(side="left")

        sel = ctk.CTkFrame(hdr, fg_color=C["card"], corner_radius=8,
                           border_width=1, border_color=C["border"])
        sel.grid(row=0, column=2, sticky="e")

        m_var = ctk.StringVar(value=MONTHS_FR[self.sel_month-1])
        y_var = ctk.StringVar(value=str(self.sel_year))

        def refresh_month(*_):
            self.sel_month = MONTHS_FR.index(m_var.get()) + 1
            self.sel_year  = int(y_var.get())
            self._go("dashboard")

        ctk.CTkOptionMenu(sel, values=MONTHS_FR, variable=m_var, command=refresh_month,
                          width=130, height=32).pack(side="left", padx=6, pady=5)
        ctk.CTkOptionMenu(sel, values=[str(y) for y in range(2020, datetime.now().year+2)],
                          variable=y_var, command=refresh_month,
                          width=88, height=32).pack(side="left", padx=(0,6), pady=5)

        # ── KPIs ──
        exp_data = self.db.get_expenses(y, m, cat_f, payee_f)
        rev_data = self.db.get_revenues(y, m)
        sav_data = self.db.get_savings(y, m)
        pat_data = self.db.get_assets(y, m)

        total_exp = sum(r["amount"] for r in exp_data)
        total_rev = sum(r["amount"] for r in rev_data)
        total_sav = sum(r["amount"] for r in sav_data)
        total_pat = sum(r["value"]  for r in pat_data)
        bilan     = total_rev - total_exp

        kpis = [
            ("Revenus",   f"{total_rev:,.0f} €", C["green"], "💶"),
            ("Dépenses",  f"{total_exp:,.0f} €", C["red"],   "💸"),
            ("Épargne",   f"{total_sav:,.0f} €", C["blue"],  "🏦"),
            ("Bilan",     f"{bilan:+,.0f} €",
             C["green"] if bilan>=0 else C["red"], "⚖️"),
        ]
        for col, (t, v, clr, ico) in enumerate(kpis):
            kpi_card(scroll, t, v, clr, ico).grid(row=1, column=col, padx=6, pady=6, sticky="ew")

        # ── Barre de filtres ──
        cats   = ["Toutes catégories"] + [c["name"] for c in self.db.get_categories()]
        payees = ["Toutes enseignes"] + self.db.get_all_payees(y, m)

        f_bar = ctk.CTkFrame(scroll, fg_color=C["card"], corner_radius=10,
                             border_width=1, border_color=C["border"])
        f_bar.grid(row=2, column=0, columnspan=4, sticky="ew", pady=(6,4))
        fb_inner = ctk.CTkFrame(f_bar, fg_color="transparent")
        fb_inner.pack(anchor="w", padx=16, pady=8)

        ctk.CTkLabel(fb_inner, text="🔍 Filtrer :", font=ctk.CTkFont(size=12, weight="bold"),
                     text_color=C["text"]).pack(side="left", padx=(0,10))

        cat_var   = ctk.StringVar(value=cat_f)
        payee_var = ctk.StringVar(value=payee_f)

        def apply_filters(*_):
            self.dash_cat_filter   = cat_var.get()
            self.dash_payee_filter = payee_var.get()
            self._go("dashboard")

        filter_bar(fb_inner, "Catégorie :", cats,   cat_var,   apply_filters)
        filter_bar(fb_inner, "Enseigne :", payees, payee_var, apply_filters)

        if cat_f not in ("Toutes catégories",) or payee_f not in ("Toutes enseignes",):
            ctk.CTkButton(fb_inner, text="✕ Effacer filtres", height=28, width=120,
                          fg_color="#FEE2E2", text_color="#EF4444", hover_color="#FECACA",
                          command=lambda: [
                              setattr(self, "dash_cat_filter", "Toutes catégories"),
                              setattr(self, "dash_payee_filter", "Toutes enseignes"),
                              self._go("dashboard")
                          ]).pack(side="left")

        # ── Graphiques côte à côte ──
        charts = ctk.CTkFrame(scroll, fg_color="transparent")
        charts.grid(row=3, column=0, columnspan=4, sticky="ew", pady=(8,0))
        charts.grid_columnconfigure((0,1), weight=1)

        # Camembert par catégorie (filtré par enseigne)
        pie_card = make_card(charts)
        pie_card.grid(row=0, column=0, padx=(0,8), sticky="nsew")
        ctk.CTkLabel(pie_card, text="Dépenses par catégorie",
                     font=ctk.CTkFont(size=14, weight="bold"), text_color=C["text"]).pack(
            anchor="w", padx=16, pady=(14,4))

        cat_data = self.db.get_expenses_by_cat(y, m, payee_f
                       if payee_f not in ("Toutes enseignes",) else None)
        if cat_data:
            fig, ax = plt.subplots(figsize=(5, 3.8))
            fig.patch.set_facecolor("white")
            labels = [r["name"] for r in cat_data]
            vals   = [r["total"] for r in cat_data]
            colors = PALETTE[:len(labels)]
            wedges, _, autotexts = ax.pie(vals, autopct="%1.0f%%", colors=colors,
                startangle=90, pctdistance=0.78, wedgeprops=dict(width=0.6))
            for at in autotexts:
                at.set_fontsize(8)
            leg_labels = [f"{l}  {v:,.0f} €" for l, v in zip(labels, vals)]
            ax.legend(wedges, leg_labels, loc="lower left",
                      bbox_to_anchor=(-0.05, -0.38), fontsize=7.5, frameon=False)
            plt.tight_layout(pad=1.5)
            cv = FigureCanvasTkAgg(fig, pie_card)
            cv.draw()
            cv.get_tk_widget().pack(fill="both", expand=True, padx=10, pady=(0,14))
            plt.close(fig)
        else:
            ctk.CTkLabel(pie_card, text="Aucune dépense ce mois",
                         text_color=C["muted"]).pack(expand=True, pady=55)

        # Top enseignes (filtré par catégorie)
        bar_card = make_card(charts)
        bar_card.grid(row=0, column=1, padx=(8,0), sticky="nsew")
        ctk.CTkLabel(bar_card, text="Top enseignes",
                     font=ctk.CTkFont(size=14, weight="bold"), text_color=C["text"]).pack(
            anchor="w", padx=16, pady=(14,4))

        payee_data = self.db.get_expenses_by_payee(y, m,
                         cat_f if cat_f not in ("Toutes catégories",) else None)
        if payee_data:
            fig2, ax2 = plt.subplots(figsize=(5, 3.8))
            fig2.patch.set_facecolor("white")
            ax2.set_facecolor("#FAFCFF")
            names = [r["payee"] for r in payee_data]
            amounts = [r["total"] for r in payee_data]
            y_pos = range(len(names))
            bars = ax2.barh(list(y_pos), amounts, color=PALETTE[0], alpha=0.85)
            ax2.set_yticks(list(y_pos))
            ax2.set_yticklabels(names, fontsize=9)
            ax2.invert_yaxis()
            ax2.set_xlabel("€", fontsize=9)
            for bar, val in zip(bars, amounts):
                ax2.text(bar.get_width() + max(amounts)*0.01, bar.get_y() + bar.get_height()/2,
                         f"{val:,.0f} €", va="center", fontsize=8)
            for sp in ["top","right"]:
                ax2.spines[sp].set_visible(False)
            plt.tight_layout(pad=1.5)
            cv2 = FigureCanvasTkAgg(fig2, bar_card)
            cv2.draw()
            cv2.get_tk_widget().pack(fill="both", expand=True, padx=10, pady=(0,14))
            plt.close(fig2)
        else:
            ctk.CTkLabel(bar_card, text="Aucune enseigne enregistrée ce mois",
                         text_color=C["muted"]).pack(expand=True, pady=55)

        # ── Patrimoine mini ──
        if pat_data:
            pc = make_card(scroll)
            pc.grid(row=4, column=0, columnspan=4, sticky="ew", pady=(14,0))
            ph = ctk.CTkFrame(pc, fg_color="transparent")
            ph.pack(fill="x", padx=16, pady=(14,6))
            ctk.CTkLabel(ph, text="📊  Patrimoine ce mois",
                         font=ctk.CTkFont(size=14, weight="bold"), text_color=C["text"]).pack(side="left")
            ctk.CTkLabel(ph, text=f"Total : {total_pat:,.0f} €",
                         font=ctk.CTkFont(size=16, weight="bold"), text_color=C["primary"]).pack(side="right")
            pin = ctk.CTkFrame(pc, fg_color="transparent")
            pin.pack(fill="x", padx=16, pady=(0,14))
            for i, bt in enumerate(self.db.get_assets_by_type(y, m)):
                pin.grid_columnconfigure(i, weight=1)
                tf = ctk.CTkFrame(pin, fg_color=C["light"], corner_radius=10)
                tf.grid(row=0, column=i, padx=5, pady=4, sticky="ew")
                ctk.CTkLabel(tf, text=ASSET_LABEL.get(bt["asset_type"], bt["asset_type"]),
                             font=ctk.CTkFont(size=10), text_color=C["muted"]).pack(anchor="w", padx=12, pady=(8,1))
                ctk.CTkLabel(tf, text=f"{bt['total']:,.0f} €",
                             font=ctk.CTkFont(size=15, weight="bold"), text_color=C["text"]).pack(
                    anchor="w", padx=12, pady=(0,8))

    # ─── PAGE: SAISIR LE MOIS ───
    def _page_monthly(self):
        top = ctk.CTkFrame(self._content, fg_color="transparent")
        top.grid(row=0, column=0, sticky="ew", padx=24, pady=(20,6))
        top.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(top, text="📝  Saisir le mois",
                     font=ctk.CTkFont(size=24, weight="bold"), text_color=C["text"]).grid(
            row=0, column=0, sticky="w")

        sel = ctk.CTkFrame(top, fg_color=C["card"], corner_radius=8,
                           border_width=1, border_color=C["border"])
        sel.grid(row=0, column=2, sticky="e")
        m_var = ctk.StringVar(value=MONTHS_FR[self.sel_month-1])
        y_var = ctk.StringVar(value=str(self.sel_year))

        def refresh(*_):
            self.sel_month = MONTHS_FR.index(m_var.get()) + 1
            self.sel_year  = int(y_var.get())
            self._go("monthly")

        ctk.CTkOptionMenu(sel, values=MONTHS_FR, variable=m_var,
                          command=refresh, width=135, height=34).pack(side="left", padx=6, pady=6)
        ctk.CTkOptionMenu(sel, values=[str(y) for y in range(2022, datetime.now().year+2)],
                          variable=y_var, command=refresh,
                          width=92, height=34).pack(side="left", padx=(0,6), pady=6)

        tabs = ctk.CTkTabview(self._content, fg_color=C["card"], corner_radius=12,
                              segmented_button_fg_color=C["light"],
                              segmented_button_selected_color=C["primary"],
                              segmented_button_selected_hover_color=C["primary"])
        tabs.grid(row=1, column=0, sticky="nsew", padx=24, pady=(0,16))
        self._content.grid_rowconfigure(1, weight=1)

        for name in ("💶  Revenus", "💸  Dépenses", "🏦  Épargne"):
            tabs.add(name)

        self._tab_revenues(tabs.tab("💶  Revenus"))
        self._tab_expenses(tabs.tab("💸  Dépenses"))
        self._tab_savings(tabs.tab("🏦  Épargne"))

    def _entry_form(self, parent, fields, on_submit, btn_label="+ Ajouter"):
        f = ctk.CTkFrame(parent, fg_color=C["light"], corner_radius=10,
                         border_width=1, border_color=C["border"])
        f.pack(fill="x", padx=2, pady=(2,8))
        entries = []
        for i, (label, ph, _) in enumerate(fields):
            f.grid_columnconfigure(i, weight=1)
            ctk.CTkLabel(f, text=label, font=ctk.CTkFont(size=11, weight="bold"),
                         text_color=C["muted"]).grid(row=0, column=i, padx=10, pady=(10,2), sticky="w")
            e = ctk.CTkEntry(f, placeholder_text=ph, height=34)
            e.grid(row=1, column=i, padx=10, pady=(0,10), sticky="ew")
            entries.append(e)
        f.grid_columnconfigure(len(fields), weight=0)
        ctk.CTkButton(f, text=btn_label, height=34, width=110,
                      command=lambda: on_submit([e.get().strip() for e in entries], entries)).grid(
            row=1, column=len(fields), padx=10, pady=(0,10))
        return entries

    def _tab_revenues(self, tab):
        wrap = ctk.CTkFrame(tab, fg_color="transparent")
        wrap.pack(fill="both", expand=True)

        data  = self.db.get_revenues(self.sel_year, self.sel_month)
        total = sum(r["amount"] for r in data)

        def on_add(vals, _):
            source, amt_str, label = vals[0], vals[1], vals[2]
            if not source or not amt_str:
                return
            try:
                amt = float(amt_str.replace(",", "."))
            except ValueError:
                return
            self.db.add_revenue(self.sel_year, self.sel_month, source, amt, label)
            self._go("monthly")

        self._entry_form(wrap,
            [("Source", "Ex : Salaire", 0), ("Montant (€)", "0.00", 0), ("Description", "Optionnel", 0)],
            on_add)

        list_f = ctk.CTkScrollableFrame(wrap, fg_color="transparent")
        list_f.pack(fill="both", expand=True)
        list_f.grid_columnconfigure((0,1,2), weight=1)

        if data:
            for col, (txt, clr) in enumerate([("Source",C["muted"]),("Montant",C["muted"]),("Description",C["muted"])]):
                ctk.CTkLabel(list_f, text=txt, font=ctk.CTkFont(size=10, weight="bold"),
                             text_color=clr).grid(row=0, column=col, padx=12, pady=(4,6), sticky="w")
            for idx, r in enumerate(data):
                add_row_widget(list_f, idx+1,
                    [(1,r["source"],C["text"]),(1,f"{r['amount']:,.2f} €",C["green"]),(2,r["label"] or "—",C["muted"])],
                    lambda rid=r["id"]: [self.db.del_revenue(rid), self._go("monthly")])
        else:
            ctk.CTkLabel(list_f, text="Aucun revenu saisi pour ce mois.",
                         text_color=C["muted"]).grid(row=0, column=0, pady=35)

        self._total_bar_pack(wrap, f"Total revenus : {total:,.2f} €", C["green"], "#F0FDF4", "#86EFAC")

    def _tab_expenses(self, tab):
        wrap = ctk.CTkFrame(tab, fg_color="transparent")
        wrap.pack(fill="both", expand=True)

        cats      = self.db.get_categories()
        cat_map   = {c["name"]: c["id"] for c in cats}
        cat_names = [c["name"] for c in cats]
        data      = self.db.get_expenses(self.sel_year, self.sel_month)
        total     = sum(r["amount"] for r in data)

        # ── Formulaire avec catégorie + enseigne ──
        form_f = ctk.CTkFrame(wrap, fg_color=C["light"], corner_radius=10,
                              border_width=1, border_color=C["border"])
        form_f.pack(fill="x", padx=2, pady=(2,8))
        form_f.grid_columnconfigure((0,1,2,3,4), weight=1)

        for i, lbl in enumerate(["Catégorie", "Enseigne / À qui", "Montant (€)", "Description"]):
            ctk.CTkLabel(form_f, text=lbl, font=ctk.CTkFont(size=11, weight="bold"),
                         text_color=C["muted"]).grid(row=0, column=i, padx=10, pady=(10,2), sticky="w")

        cat_var = ctk.StringVar(value=cat_names[0] if cat_names else "")
        ctk.CTkOptionMenu(form_f, values=cat_names, variable=cat_var,
                          height=34).grid(row=1, column=0, padx=10, pady=(0,10), sticky="ew")

        payee_e = ctk.CTkEntry(form_f, placeholder_text="Ex : Carrefour", height=34)
        payee_e.grid(row=1, column=1, padx=10, pady=(0,10), sticky="ew")

        amt_e = ctk.CTkEntry(form_f, placeholder_text="0.00", height=34)
        amt_e.grid(row=1, column=2, padx=10, pady=(0,10), sticky="ew")

        desc_e = ctk.CTkEntry(form_f, placeholder_text="Optionnel", height=34)
        desc_e.grid(row=1, column=3, padx=10, pady=(0,10), sticky="ew")

        def on_add():
            cat_name = cat_var.get()
            payee    = payee_e.get().strip()
            try:
                amt = float(amt_e.get().strip().replace(",", "."))
            except ValueError:
                return
            cid = cat_map.get(cat_name)
            if not cid:
                return
            self.db.add_expense(self.sel_year, self.sel_month, cid, amt,
                                 desc_e.get().strip(), payee)
            self._go("monthly")

        ctk.CTkButton(form_f, text="+ Ajouter", height=34, width=100,
                      command=on_add).grid(row=1, column=4, padx=10, pady=(0,10))

        # ── Liste ──
        list_f = ctk.CTkScrollableFrame(wrap, fg_color="transparent")
        list_f.pack(fill="both", expand=True)
        list_f.grid_columnconfigure((0,1,2,3), weight=1)

        if data:
            for col, txt in enumerate(["Catégorie","Enseigne","Montant","Description"]):
                ctk.CTkLabel(list_f, text=txt, font=ctk.CTkFont(size=10, weight="bold"),
                             text_color=C["muted"]).grid(row=0, column=col, padx=12, pady=(4,6), sticky="w")
            for idx, r in enumerate(data):
                add_row_widget(list_f, idx+1,
                    [(1,r["cat"],C["text"]),
                     (1,r["payee"] or "—",C["muted"]),
                     (1,f"{r['amount']:,.2f} €",C["red"]),
                     (2,r["label"] or "—",C["muted"])],
                    lambda eid=r["id"]: [self.db.del_expense(eid), self._go("monthly")])
        else:
            ctk.CTkLabel(list_f, text="Aucune dépense saisie pour ce mois.",
                         text_color=C["muted"]).grid(row=0, column=0, pady=35)

        self._total_bar_pack(wrap, f"Total dépenses : {total:,.2f} €", C["red"], "#FEF2F2", "#FCA5A5")

    def _tab_savings(self, tab):
        wrap = ctk.CTkFrame(tab, fg_color="transparent")
        wrap.pack(fill="both", expand=True)

        data  = self.db.get_savings(self.sel_year, self.sel_month)
        total = sum(r["amount"] for r in data)

        def on_add(vals, _):
            account, amt_str, label = vals[0], vals[1], vals[2]
            if not account or not amt_str:
                return
            try:
                amt = float(amt_str.replace(",", "."))
            except ValueError:
                return
            self.db.add_saving(self.sel_year, self.sel_month, account, amt, label)
            self._go("monthly")

        self._entry_form(wrap,
            [("Compte / Destination","Ex : Livret A",0),("Montant (€)","0.00",0),("Description","Optionnel",0)],
            on_add)

        list_f = ctk.CTkScrollableFrame(wrap, fg_color="transparent")
        list_f.pack(fill="both", expand=True)
        list_f.grid_columnconfigure((0,1,2), weight=1)

        if data:
            for col, txt in enumerate(["Compte","Montant","Description"]):
                ctk.CTkLabel(list_f, text=txt, font=ctk.CTkFont(size=10, weight="bold"),
                             text_color=C["muted"]).grid(row=0, column=col, padx=12, pady=(4,6), sticky="w")
            for idx, r in enumerate(data):
                add_row_widget(list_f, idx+1,
                    [(1,r["account"] or "—",C["text"]),
                     (1,f"{r['amount']:,.2f} €",C["blue"]),
                     (2,r["label"] or "—",C["muted"])],
                    lambda sid=r["id"]: [self.db.del_saving(sid), self._go("monthly")])
        else:
            ctk.CTkLabel(list_f, text="Aucune épargne saisie pour ce mois.",
                         text_color=C["muted"]).grid(row=0, column=0, pady=35)

        self._total_bar_pack(wrap, f"Total épargne : {total:,.2f} €", C["blue"], "#EFF6FF", "#93C5FD")

    def _total_bar(self, parent, text, color, bg, border, row=2):
        f = ctk.CTkFrame(parent, fg_color=bg, corner_radius=8, border_width=1, border_color=border)
        f.grid(row=row, column=0, sticky="ew", padx=2, pady=(6,4))
        ctk.CTkLabel(f, text=text, font=ctk.CTkFont(size=14, weight="bold"),
                     text_color=color).pack(side="right", padx=16, pady=9)

    def _total_bar_pack(self, parent, text, color, bg, border):
        f = ctk.CTkFrame(parent, fg_color=bg, corner_radius=8, border_width=1, border_color=border)
        f.pack(fill="x", padx=2, pady=(6,4))
        ctk.CTkLabel(f, text=text, font=ctk.CTkFont(size=14, weight="bold"),
                     text_color=color).pack(side="right", padx=16, pady=9)

    # ─── PAGE: ANALYSES ───
    def _page_analyses(self):
        y, m = self.sel_year, self.sel_month

        ctk.CTkLabel(self._content, text="📊  Analyses détaillées",
                     font=ctk.CTkFont(size=24, weight="bold"), text_color=C["text"]).grid(
            row=0, column=0, sticky="w", padx=24, pady=(20,6))

        # Sélecteur mois
        sel = ctk.CTkFrame(self._content, fg_color="transparent")
        sel.grid(row=1, column=0, sticky="ew", padx=24, pady=(0,6))
        m_var = ctk.StringVar(value=MONTHS_FR[self.sel_month-1])
        y_var = ctk.StringVar(value=str(self.sel_year))

        def refresh(*_):
            self.sel_month = MONTHS_FR.index(m_var.get()) + 1
            self.sel_year  = int(y_var.get())
            self._go("analyses")

        sel_box = ctk.CTkFrame(sel, fg_color=C["card"], corner_radius=8,
                               border_width=1, border_color=C["border"])
        sel_box.pack(side="left")
        ctk.CTkOptionMenu(sel_box, values=MONTHS_FR, variable=m_var,
                          command=refresh, width=135, height=32).pack(side="left", padx=6, pady=5)
        ctk.CTkOptionMenu(sel_box, values=[str(y2) for y2 in range(2022, datetime.now().year+2)],
                          variable=y_var, command=refresh,
                          width=88, height=32).pack(side="left", padx=(0,6), pady=5)

        scroll = ctk.CTkScrollableFrame(self._content, fg_color="transparent")
        scroll.grid(row=2, column=0, sticky="nsew", padx=24, pady=(0,16))
        scroll.grid_columnconfigure((0,1), weight=1)
        self._content.grid_rowconfigure(2, weight=1)

        # ── Évolution 6 mois : revenus / dépenses ──
        ev_card = make_card(scroll)
        ev_card.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0,12))
        ctk.CTkLabel(ev_card, text="Évolution revenus / dépenses sur 6 mois",
                     font=ctk.CTkFont(size=14, weight="bold"), text_color=C["text"]).pack(
            anchor="w", padx=16, pady=(14,4))

        summaries = list(reversed(self.db.monthly_summary(6)))
        if summaries:
            fig, ax = plt.subplots(figsize=(10, 3.4))
            fig.patch.set_facecolor("white")
            ax.set_facecolor("#FAFCFF")
            x       = range(len(summaries))
            xlabels = [f"{MONTHS_FR[r['month']-1][:3]}\n{r['year']}" for r in summaries]
            rev_v   = [r["rev"] for r in summaries]
            exp_v   = [r["exp"] for r in summaries]
            sav_v   = [r["sav"] for r in summaries]
            w = 0.28
            ax.bar([i-w for i in x], rev_v, w, color="#22C55E", alpha=0.85, label="Revenus")
            ax.bar([i   for i in x], exp_v, w, color="#EF4444", alpha=0.85, label="Dépenses")
            ax.bar([i+w for i in x], sav_v, w, color="#3B82F6", alpha=0.85, label="Épargne")
            ax.set_xticks(list(x))
            ax.set_xticklabels(xlabels, fontsize=9)
            ax.legend(fontsize=10, frameon=False)
            ax.set_ylabel("€", fontsize=9)
            ax.yaxis.set_tick_params(labelsize=8)
            for sp in ["top","right"]:
                ax.spines[sp].set_visible(False)
            plt.tight_layout(pad=1.5)
            cv = FigureCanvasTkAgg(fig, ev_card)
            cv.draw()
            cv.get_tk_widget().pack(fill="both", expand=True, padx=12, pady=(0,14))
            plt.close(fig)

        # ── Camembert catégories + Top enseignes ──
        cat_data   = self.db.get_expenses_by_cat(y, m)
        payee_data = self.db.get_expenses_by_payee(y, m)

        pie2_card = make_card(scroll)
        pie2_card.grid(row=1, column=0, padx=(0,8), sticky="nsew", pady=(0,12))
        ctk.CTkLabel(pie2_card, text=f"Catégories — {MONTHS_FR[m-1]} {y}",
                     font=ctk.CTkFont(size=14, weight="bold"), text_color=C["text"]).pack(
            anchor="w", padx=16, pady=(14,4))

        if cat_data:
            fig3, ax3 = plt.subplots(figsize=(5, 4.2))
            fig3.patch.set_facecolor("white")
            labels3 = [r["name"] for r in cat_data]
            vals3   = [r["total"] for r in cat_data]
            colors3 = PALETTE[:len(labels3)]
            wedges, _, ats = ax3.pie(vals3, autopct="%1.0f%%", colors=colors3,
                startangle=90, pctdistance=0.78, wedgeprops=dict(width=0.6))
            for at in ats:
                at.set_fontsize(8)
            ax3.legend(wedges, [f"{l}  {v:,.0f} €" for l,v in zip(labels3,vals3)],
                       loc="lower left", bbox_to_anchor=(-0.05,-0.4), fontsize=7.5, frameon=False)
            plt.tight_layout(pad=1.5)
            cv3 = FigureCanvasTkAgg(fig3, pie2_card)
            cv3.draw()
            cv3.get_tk_widget().pack(fill="both", expand=True, padx=10, pady=(0,14))
            plt.close(fig3)
        else:
            ctk.CTkLabel(pie2_card, text="Aucune dépense ce mois",
                         text_color=C["muted"]).pack(expand=True, pady=55)

        bar2_card = make_card(scroll)
        bar2_card.grid(row=1, column=1, padx=(8,0), sticky="nsew", pady=(0,12))
        ctk.CTkLabel(bar2_card, text=f"Top enseignes — {MONTHS_FR[m-1]} {y}",
                     font=ctk.CTkFont(size=14, weight="bold"), text_color=C["text"]).pack(
            anchor="w", padx=16, pady=(14,4))

        if payee_data:
            fig4, ax4 = plt.subplots(figsize=(5, 4.2))
            fig4.patch.set_facecolor("white")
            ax4.set_facecolor("#FAFCFF")
            pnames = [r["payee"] for r in payee_data]
            pamounts = [r["total"] for r in payee_data]
            yp = range(len(pnames))
            ax4.barh(list(yp), pamounts, color=PALETTE[:len(pnames)], alpha=0.85)
            ax4.set_yticks(list(yp))
            ax4.set_yticklabels(pnames, fontsize=8)
            ax4.invert_yaxis()
            ax4.set_xlabel("€", fontsize=9)
            for sp in ["top","right"]:
                ax4.spines[sp].set_visible(False)
            plt.tight_layout(pad=1.5)
            cv4 = FigureCanvasTkAgg(fig4, bar2_card)
            cv4.draw()
            cv4.get_tk_widget().pack(fill="both", expand=True, padx=10, pady=(0,14))
            plt.close(fig4)
        else:
            ctk.CTkLabel(bar2_card, text="Aucune enseigne ce mois",
                         text_color=C["muted"]).pack(expand=True, pady=55)

    # ─── PAGE: PATRIMOINE ───
    def _page_patrimoine(self):
        top = ctk.CTkFrame(self._content, fg_color="transparent")
        top.grid(row=0, column=0, sticky="ew", padx=24, pady=(20,6))
        top.grid_columnconfigure(1, weight=1)
        ctk.CTkLabel(top, text="📈  Patrimoine & Investissements",
                     font=ctk.CTkFont(size=24, weight="bold"), text_color=C["text"]).grid(
            row=0, column=0, sticky="w")

        sel = ctk.CTkFrame(top, fg_color=C["card"], corner_radius=8,
                           border_width=1, border_color=C["border"])
        sel.grid(row=0, column=2, sticky="e")
        m_var = ctk.StringVar(value=MONTHS_FR[self.sel_month-1])
        y_var = ctk.StringVar(value=str(self.sel_year))

        def refresh(*_):
            self.sel_month = MONTHS_FR.index(m_var.get()) + 1
            self.sel_year  = int(y_var.get())
            self._go("patrimoine")

        ctk.CTkOptionMenu(sel, values=MONTHS_FR, variable=m_var,
                          command=refresh, width=135, height=34).pack(side="left", padx=6, pady=6)
        ctk.CTkOptionMenu(sel, values=[str(y) for y in range(2022, datetime.now().year+2)],
                          variable=y_var, command=refresh,
                          width=92, height=34).pack(side="left", padx=(0,6), pady=6)

        scroll = ctk.CTkScrollableFrame(self._content, fg_color="transparent")
        scroll.grid(row=1, column=0, sticky="nsew", padx=24, pady=(0,16))
        scroll.grid_columnconfigure((0,1), weight=1)
        self._content.grid_rowconfigure(1, weight=1)

        # Formulaire ajout
        fc = make_card(scroll)
        fc.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0,12))
        ctk.CTkLabel(fc, text="Ajouter / Mettre à jour un actif",
                     font=ctk.CTkFont(size=14, weight="bold"), text_color=C["text"]).pack(
            anchor="w", padx=16, pady=(14,6))

        fi = ctk.CTkFrame(fc, fg_color=C["light"], corner_radius=8)
        fi.pack(fill="x", padx=16, pady=(0,14))
        fi.grid_columnconfigure((0,1,2,3,4), weight=1)

        for i, lbl in enumerate(["Type","Nom","Valeur (€)","Notes"]):
            ctk.CTkLabel(fi, text=lbl, font=ctk.CTkFont(size=11, weight="bold"),
                         text_color=C["muted"]).grid(row=0, column=i, padx=10, pady=(10,2), sticky="w")

        type_names = [t[0] for t in ASSET_TYPES]
        type_keys  = {t[0]: t[1] for t in ASSET_TYPES}
        type_var   = ctk.StringVar(value=type_names[0])
        ctk.CTkOptionMenu(fi, values=type_names, variable=type_var,
                          height=34).grid(row=1, column=0, padx=10, pady=(0,12), sticky="ew")
        name_e  = ctk.CTkEntry(fi, placeholder_text="Ex : PEA Boursorama", height=34)
        name_e.grid(row=1, column=1, padx=10, pady=(0,12), sticky="ew")
        val_e   = ctk.CTkEntry(fi, placeholder_text="0.00", height=34)
        val_e.grid(row=1, column=2, padx=10, pady=(0,12), sticky="ew")
        notes_e = ctk.CTkEntry(fi, placeholder_text="Optionnel", height=34)
        notes_e.grid(row=1, column=3, padx=10, pady=(0,12), sticky="ew")

        def save_asset():
            name = name_e.get().strip()
            try:
                val = float(val_e.get().strip().replace(",", "."))
            except ValueError:
                return
            if not name:
                return
            self.db.upsert_asset(self.sel_year, self.sel_month,
                                  type_keys.get(type_var.get(), "autres"),
                                  name, val, notes_e.get().strip())
            self._go("patrimoine")

        ctk.CTkButton(fi, text="✓  Sauvegarder", height=34,
                      command=save_asset).grid(row=1, column=4, padx=10, pady=(0,12))

        # Liste actifs
        assets = self.db.get_assets(self.sel_year, self.sel_month)
        total  = sum(a["value"] for a in assets)

        lc = make_card(scroll)
        lc.grid(row=1, column=0, padx=(0,8), sticky="nsew")
        ctk.CTkLabel(lc, text=f"Actifs — {MONTHS_FR[self.sel_month-1]} {self.sel_year}",
                     font=ctk.CTkFont(size=14, weight="bold"), text_color=C["text"]).pack(
            anchor="w", padx=16, pady=(14,6))

        if assets:
            cur_type = None
            for a in assets:
                if a["asset_type"] != cur_type:
                    cur_type = a["asset_type"]
                    ctk.CTkLabel(lc, text=f"  {ASSET_LABEL.get(cur_type,cur_type)}",
                                 font=ctk.CTkFont(size=11, weight="bold"),
                                 text_color=C["muted"]).pack(anchor="w", padx=10, pady=(6,2))
                rf = ctk.CTkFrame(lc, fg_color=C["light"], corner_radius=6)
                rf.pack(fill="x", padx=12, pady=2)
                rf.grid_columnconfigure(1, weight=1)
                ctk.CTkLabel(rf, text=a["asset_name"],
                             font=ctk.CTkFont(size=12)).grid(row=0, column=0, padx=12, pady=8, sticky="w")
                ctk.CTkLabel(rf, text=f"{a['value']:,.2f} €",
                             font=ctk.CTkFont(size=12, weight="bold"),
                             text_color=C["primary"]).grid(row=0, column=1, padx=12, pady=8, sticky="e")
                def del_a(aid=a["id"]):
                    self.db.del_asset(aid); self._go("patrimoine")
                ctk.CTkButton(rf, text="✕", width=28, height=26, fg_color="#FEE2E2",
                              text_color="#EF4444", hover_color="#FECACA",
                              command=del_a).grid(row=0, column=2, padx=8, pady=4)
            tf = ctk.CTkFrame(lc, fg_color="#EFF6FF", corner_radius=8,
                              border_width=1, border_color="#93C5FD")
            tf.pack(fill="x", padx=12, pady=(10,14))
            ctk.CTkLabel(tf, text=f"Total : {total:,.2f} €",
                         font=ctk.CTkFont(size=15, weight="bold"),
                         text_color=C["primary"]).pack(side="right", padx=16, pady=9)
        else:
            ctk.CTkLabel(lc, text="Aucun actif saisi pour ce mois.",
                         text_color=C["muted"]).pack(expand=True, pady=40)

        # Graphique évolution
        rc = make_card(scroll)
        rc.grid(row=1, column=1, padx=(8,0), sticky="nsew")
        ctk.CTkLabel(rc, text="Évolution du patrimoine",
                     font=ctk.CTkFont(size=14, weight="bold"), text_color=C["text"]).pack(
            anchor="w", padx=16, pady=(14,4))

        history = self.db.get_asset_history()
        if len(history) >= 2:
            fig, ax = plt.subplots(figsize=(5.5, 4))
            fig.patch.set_facecolor("white")
            ax.set_facecolor("#FAFCFF")
            pts    = [(r["year"],r["month"],r["total"]) for r in history]
            labels = [f"{MONTHS_FR[p[1]-1][:3]}\n{p[0]}" for p in pts]
            vals   = [p[2] for p in pts]
            ax.plot(range(len(vals)), vals, color=C["primary"], linewidth=2.5,
                    marker="o", markersize=5, zorder=3)
            ax.fill_between(range(len(vals)), vals, alpha=0.08, color=C["primary"])
            ax.set_xticks(range(len(labels)))
            ax.set_xticklabels(labels, fontsize=7.5)
            ax.yaxis.set_tick_params(labelsize=8)
            ax.set_ylabel("€", fontsize=9)
            for sp in ["top","right"]:
                ax.spines[sp].set_visible(False)
            plt.tight_layout(pad=1.5)
            cv = FigureCanvasTkAgg(fig, rc)
            cv.draw()
            cv.get_tk_widget().pack(fill="both", expand=True, padx=10, pady=(0,14))
            plt.close(fig)
        else:
            ctk.CTkLabel(rc, text="Saisissez des actifs sur\nplusieurs mois pour voir\nl'évolution.",
                         text_color=C["muted"], justify="center").pack(expand=True, pady=55)

    # ─── PAGE: HISTORIQUE ───
    def _page_historique(self):
        ctk.CTkLabel(self._content, text="📋  Historique",
                     font=ctk.CTkFont(size=24, weight="bold"), text_color=C["text"]).grid(
            row=0, column=0, sticky="w", padx=24, pady=(20,6))

        scroll = ctk.CTkScrollableFrame(self._content, fg_color="transparent")
        scroll.grid(row=1, column=0, sticky="nsew", padx=24, pady=(0,16))
        scroll.grid_columnconfigure(0, weight=1)
        self._content.grid_rowconfigure(1, weight=1)

        rows = self.db.monthly_summary(48)

        if not rows:
            ctk.CTkLabel(scroll, text="Aucune donnée. Commencez par 'Saisir le mois' !",
                         text_color=C["muted"], font=ctk.CTkFont(size=14)).grid(row=0,column=0,pady=60)
            return

        hdr = ctk.CTkFrame(scroll, fg_color=C["sidebar"], corner_radius=8)
        hdr.grid(row=0, column=0, sticky="ew", pady=(0,4))
        hdr.grid_columnconfigure((0,1,2,3,4), weight=1)
        for i, h in enumerate(["Mois","Revenus","Dépenses","Épargne","Bilan"]):
            ctk.CTkLabel(hdr, text=h, font=ctk.CTkFont(size=12, weight="bold"),
                         text_color="white").grid(row=0, column=i, padx=16, pady=10, sticky="w")

        for idx, r in enumerate(rows):
            bilan = r["rev"] - r["exp"]
            bg = C["card"] if idx%2==0 else C["light"]
            rf = ctk.CTkFrame(scroll, fg_color=bg, corner_radius=6)
            rf.grid(row=idx+1, column=0, sticky="ew", pady=1)
            rf.grid_columnconfigure((0,1,2,3,4), weight=1)
            ctk.CTkLabel(rf, text=f"{MONTHS_FR[r['month']-1]} {r['year']}",
                         font=ctk.CTkFont(size=13, weight="bold"), text_color=C["text"]).grid(
                row=0, column=0, padx=16, pady=10, sticky="w")
            ctk.CTkLabel(rf, text=f"{r['rev']:,.0f} €",
                         text_color=C["green"], font=ctk.CTkFont(size=13)).grid(
                row=0, column=1, padx=16, pady=10, sticky="w")
            ctk.CTkLabel(rf, text=f"{r['exp']:,.0f} €",
                         text_color=C["red"], font=ctk.CTkFont(size=13)).grid(
                row=0, column=2, padx=16, pady=10, sticky="w")
            ctk.CTkLabel(rf, text=f"{r['sav']:,.0f} €",
                         text_color=C["blue"], font=ctk.CTkFont(size=13)).grid(
                row=0, column=3, padx=16, pady=10, sticky="w")
            ctk.CTkLabel(rf, text=f"{bilan:+,.0f} €",
                         text_color=C["green"] if bilan>=0 else C["red"],
                         font=ctk.CTkFont(size=13, weight="bold")).grid(
                row=0, column=4, padx=16, pady=10, sticky="w")

        ef = ctk.CTkFrame(scroll, fg_color="transparent")
        ef.grid(row=len(rows)+1, column=0, sticky="e", pady=14)
        ctk.CTkButton(ef, text="📄  Exporter CSV", command=self._export_csv, height=36,
                      fg_color=C["muted"], hover_color="#475569").pack(side="right", padx=(6,0))
        ctk.CTkButton(ef, text="📊  Exporter Excel", command=self._export_excel, height=36,
                      fg_color="#16A34A", hover_color="#15803D").pack(side="right")

    def _export_csv(self):
        path = fd.asksaveasfilename(defaultextension=".csv",
            filetypes=[("CSV","*.csv")], initialfile="finance_tracker.csv")
        if path:
            self.db.export_csv(path)
            self._toast("✅  Export CSV réussi !")

    def _export_excel(self):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            import subprocess, sys
            subprocess.run([sys.executable,"-m","pip","install","openpyxl"], capture_output=True)
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment

        path = fd.asksaveasfilename(defaultextension=".xlsx",
            filetypes=[("Excel","*.xlsx")], initialfile="finance_tracker.xlsx")
        if not path:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Résumé mensuel"
        hfill = PatternFill("solid", fgColor="1A2340")
        hfont = Font(bold=True, color="FFFFFF", size=11)
        headers = ["Année","Mois","Revenus (€)","Dépenses (€)","Épargne (€)","Bilan (€)"]
        for ci, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=ci, value=h)
            cell.fill = hfill; cell.font = hfont
            cell.alignment = Alignment(horizontal="center")
        for ri, r in enumerate(reversed(self.db.monthly_summary(120)), 2):
            bilan = r["rev"] - r["exp"]
            ws.cell(ri,1,r["year"]); ws.cell(ri,2,MONTHS_FR[r["month"]-1])
            ws.cell(ri,3,round(r["rev"],2)); ws.cell(ri,4,round(r["exp"],2))
            ws.cell(ri,5,round(r["sav"],2))
            c = ws.cell(ri,6,round(bilan,2))
            c.font = Font(color="16A34A" if bilan>=0 else "EF4444", bold=True)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = \
                max(len(str(cell.value or "")) for cell in col) + 4
        wb.save(path)
        self._toast("✅  Export Excel réussi !")

    # ─── PAGE: PARAMÈTRES ───
    def _page_settings(self):
        ctk.CTkLabel(self._content, text="⚙️  Paramètres",
                     font=ctk.CTkFont(size=24, weight="bold"), text_color=C["text"]).grid(
            row=0, column=0, sticky="w", padx=24, pady=(20,6))

        scroll = ctk.CTkScrollableFrame(self._content, fg_color="transparent")
        scroll.grid(row=1, column=0, sticky="nsew", padx=24, pady=(0,16))
        scroll.grid_columnconfigure(0, weight=1)
        self._content.grid_rowconfigure(1, weight=1)

        # ── Import Notion ──
        imp_card = make_card(scroll)
        imp_card.grid(row=0, column=0, sticky="ew", pady=(0,14))
        ctk.CTkLabel(imp_card, text="📥  Importer depuis Notion (ZIP)",
                     font=ctk.CTkFont(size=14, weight="bold"), text_color=C["text"]).pack(
            anchor="w", padx=16, pady=(14,4))
        ctk.CTkLabel(imp_card,
            text="Importe directement les exports ZIP de Notion (revenus ou dépenses).",
            font=ctk.CTkFont(size=12), text_color=C["muted"]).pack(anchor="w", padx=16, pady=(0,6))

        ib = ctk.CTkFrame(imp_card, fg_color="transparent")
        ib.pack(anchor="w", padx=16, pady=(0,14))
        ctk.CTkButton(ib, text="🗜  Importer ZIP Revenus",
                      command=lambda: self._import_notion("revenus"),
                      height=36, fg_color=C["green"], hover_color="#15803D").pack(side="left", padx=(0,8))
        ctk.CTkButton(ib, text="🗜  Importer ZIP Dépenses",
                      command=lambda: self._import_notion("dépenses"),
                      height=36, fg_color=C["primary"]).pack(side="left", padx=(0,8))
        ctk.CTkButton(ib, text="📄  Importer CSV",
                      command=self._import_csv_legacy,
                      height=36, fg_color=C["muted"], hover_color="#475569").pack(side="left")

        # ── Catégories ──
        cat_card = make_card(scroll)
        cat_card.grid(row=1, column=0, sticky="ew", pady=(0,14))
        ctk.CTkLabel(cat_card, text="📂  Catégories de dépenses",
                     font=ctk.CTkFont(size=14, weight="bold"), text_color=C["text"]).pack(
            anchor="w", padx=16, pady=(14,8))

        add_f = ctk.CTkFrame(cat_card, fg_color=C["light"], corner_radius=8)
        add_f.pack(fill="x", padx=16, pady=(0,10))
        new_e = ctk.CTkEntry(add_f, placeholder_text="Nouvelle catégorie…", height=34, width=280)
        new_e.pack(side="left", padx=10, pady=10)
        ctk.CTkButton(add_f, text="+ Ajouter", height=34, width=110,
                      command=lambda: [self.db.add_category(new_e.get().strip()), self._go("settings")]
                      ).pack(side="left", padx=(0,10), pady=10)

        grid_f = ctk.CTkFrame(cat_card, fg_color="transparent")
        grid_f.pack(fill="x", padx=16, pady=(0,14))
        cats = self.db.get_categories()
        cols = 3
        for i, cat in enumerate(cats):
            r, c = divmod(i, cols)
            grid_f.grid_columnconfigure(c, weight=1)
            item = ctk.CTkFrame(grid_f, fg_color=C["light"], corner_radius=8)
            item.grid(row=r, column=c, padx=4, pady=4, sticky="ew")
            ctk.CTkLabel(item, text=cat["name"],
                         font=ctk.CTkFont(size=12), text_color=C["text"]).pack(
                side="left", padx=10, pady=7)
            if not cat["is_default"]:
                def del_c(cid=cat["id"]):
                    self.db.del_category(cid); self._go("settings")
                ctk.CTkButton(item, text="✕", width=26, height=24,
                              fg_color="#FEE2E2", text_color="#EF4444",
                              hover_color="#FECACA", command=del_c).pack(side="right", padx=6, pady=4)
            else:
                ctk.CTkLabel(item, text="défaut", font=ctk.CTkFont(size=9),
                             text_color=C["muted"]).pack(side="right", padx=8)

        # ── Info ──
        info_card = make_card(scroll)
        info_card.grid(row=2, column=0, sticky="ew")
        ctk.CTkLabel(info_card, text="ℹ️  Informations",
                     font=ctk.CTkFont(size=14, weight="bold"), text_color=C["text"]).pack(
            anchor="w", padx=16, pady=(14,4))
        ctk.CTkLabel(info_card, text=f"📁  Base de données : {DB_PATH}",
                     font=ctk.CTkFont(size=12), text_color=C["muted"]).pack(
            anchor="w", padx=16, pady=(0,14))

    # ─── IMPORT ───
    def _import_notion(self, kind: str):
        path = fd.askopenfilename(
            title=f"Importer ZIP {kind}",
            filetypes=[("ZIP","*.zip"),("Tous","*.*")])
        if not path:
            return
        try:
            rev, exp = self.db.import_notion_zip(path)
            total = rev + exp
            if total > 0:
                self._toast(f"✅  {rev} revenus + {exp} dépenses importés !")
                self._go("historique")
            else:
                mb.showwarning("Import", "Aucune donnée trouvée dans ce ZIP.")
        except Exception as e:
            mb.showerror("Erreur import", str(e))

    def _import_csv_legacy(self):
        path = fd.askopenfilename(title="Importer CSV",
            filetypes=[("CSV","*.csv"),("Tous","*.*")])
        if not path:
            return
        count = 0
        try:
            with open(path, newline="", encoding="utf-8-sig") as f:
                reader = csv.DictReader(f, delimiter=";")
                for raw in reader:
                    row = {k.strip(): (v or "").strip() for k,v in raw.items()}
                    try:
                        year = int(row.get("Année", row.get("Annee", 0)))
                        mname = row.get("Mois","")
                        if mname not in MONTHS_FR:
                            continue
                        month = MONTHS_FR.index(mname)+1
                        rev = float(row.get("Revenus (€)","0").replace(",","."))
                        exp = float(row.get("Dépenses (€)","0").replace(",","."))
                        sav = float(row.get("Épargne (€)","0").replace(",","."))
                        mid = self.db._month_id(year, month, create=True)
                        if self.db.con.execute("SELECT COUNT(*) FROM revenues WHERE month_id=?",
                                               (mid,)).fetchone()[0] == 0 and rev > 0:
                            self.db.add_revenue(year, month, "Import", rev)
                        if self.db.con.execute("SELECT COUNT(*) FROM savings WHERE month_id=?",
                                               (mid,)).fetchone()[0] == 0 and sav > 0:
                            self.db.add_saving(year, month, "Import", sav)
                        count += 1
                    except Exception:
                        continue
        except Exception as e:
            mb.showerror("Erreur", str(e))
            return
        if count:
            self._toast(f"✅  {count} mois importés !")
            self._go("historique")
        else:
            mb.showwarning("Import CSV", "Aucune donnée trouvée.")

    # ─── TOAST ───
    def _toast(self, msg):
        t = ctk.CTkToplevel(self)
        t.geometry(f"320x54+{self.winfo_x()+self.winfo_width()//2-160}+{self.winfo_y()+self.winfo_height()-90}")
        t.overrideredirect(True)
        t.attributes("-topmost", True)
        ctk.CTkLabel(t, text=msg, font=ctk.CTkFont(size=13, weight="bold"),
                     fg_color=C["sidebar"], text_color="white",
                     corner_radius=10).pack(fill="both", expand=True, padx=4, pady=4)
        t.after(2800, t.destroy)


# ─────────────────────────────────────────────
if __name__ == "__main__":
    app = App()
    app.mainloop()
