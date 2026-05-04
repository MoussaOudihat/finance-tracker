"""
utils_excel.py — Export / Import Excel mensuel (openpyxl)

Export : deux feuilles (Dépenses + Revenus) pour le mois sélectionné.
Import : lit le même format et insère les lignes dans la DB.
"""
import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from logger import log


# ── Couleurs Fintrack ──────────────────────────────────────────
_PRIMARY   = "4F46E5"   # indigo
_PRIMARY_L = "EEF2FF"   # indigo très clair
_RED_H     = "EF4444"
_RED_L     = "FEF2F2"
_GREEN_H   = "22C55E"
_GREEN_L   = "F0FDF4"
_MUTED     = "64748B"
_BORDER_C  = "E2E8F0"

_THIN = Side(style="thin", color=_BORDER_C)
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _hdr_style(cell, bg: str = _PRIMARY):
    cell.font      = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
    cell.fill      = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border    = _BORDER


def _data_style(cell, bg: str = "FFFFFF", color: str = "1E293B",
                bold: bool = False, align: str = "left"):
    cell.font      = Font(name="Calibri", size=10, color=color, bold=bold)
    cell.fill      = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border    = _BORDER


def _autofit(ws, min_width: int = 12, max_width: int = 45):
    for col in ws.columns:
        width = min_width
        for cell in col:
            if cell.value:
                width = max(width, min(max_width, len(str(cell.value)) + 4))
        ws.column_dimensions[get_column_letter(col[0].column)].width = width


def _write_expenses_sheet(ws, expenses: list, month_label: str):
    ws.title = "Dépenses"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 32

    # Titre
    ws.merge_cells("A1:D1")
    title = ws["A1"]
    title.value = f"💸  Dépenses — {month_label}"
    title.font  = Font(bold=True, size=14, color=_RED_H, name="Calibri")
    title.fill  = PatternFill("solid", fgColor=_RED_L)
    title.alignment = Alignment(horizontal="center", vertical="center")

    # En-têtes
    headers = ["Catégorie", "Enseigne", "Montant (€)", "Description"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=2, column=col, value=h)
        _hdr_style(c, bg=_RED_H)
    ws.row_dimensions[2].height = 24

    # Données
    for i, row in enumerate(expenses):
        r = i + 3
        bg = "FFFFFF" if i % 2 == 0 else "FEF9FA"
        ws.row_dimensions[r].height = 20

        c_cat = ws.cell(r, 1, row["cat"])
        _data_style(c_cat, bg=bg)

        c_pay = ws.cell(r, 2, row["payee"] or "")
        _data_style(c_pay, bg=bg, color=_MUTED)

        c_amt = ws.cell(r, 3, round(row["amount"], 2))
        _data_style(c_amt, bg=bg, color=_RED_H, bold=True, align="right")
        c_amt.number_format = '#,##0.00 "€"'

        c_lbl = ws.cell(r, 4, row["label"] or "")
        _data_style(c_lbl, bg=bg, color=_MUTED)

    # Ligne total
    if expenses:
        tr = len(expenses) + 3
        ws.row_dimensions[tr].height = 22
        ws.cell(tr, 1, "TOTAL").font = Font(bold=True, size=11, name="Calibri")
        ws.cell(tr, 1).fill = PatternFill("solid", fgColor=_RED_L)
        ws.cell(tr, 1).alignment = Alignment(horizontal="right")
        ws.merge_cells(f"A{tr}:B{tr}")

        total_cell = ws.cell(tr, 3, f"=SUM(C3:C{tr-1})")
        _data_style(total_cell, bg=_RED_L, color=_RED_H, bold=True, align="right")
        total_cell.number_format = '#,##0.00 "€"'

        ws.cell(tr, 4).fill = PatternFill("solid", fgColor=_RED_L)

    # Freeze + autofit
    ws.freeze_panes = "A3"
    _autofit(ws)


def _write_revenues_sheet(ws, revenues: list, month_label: str):
    ws.title = "Revenus"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 32

    # Titre
    ws.merge_cells("A1:C1")
    title = ws["A1"]
    title.value = f"💶  Revenus — {month_label}"
    title.font  = Font(bold=True, size=14, color=_GREEN_H, name="Calibri")
    title.fill  = PatternFill("solid", fgColor=_GREEN_L)
    title.alignment = Alignment(horizontal="center", vertical="center")

    # En-têtes
    headers = ["Source", "Montant (€)", "Description"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=2, column=col, value=h)
        _hdr_style(c, bg=_GREEN_H)
    ws.row_dimensions[2].height = 24

    # Données
    for i, row in enumerate(revenues):
        r = i + 3
        bg = "FFFFFF" if i % 2 == 0 else "F0FDF9"
        ws.row_dimensions[r].height = 20

        c_src = ws.cell(r, 1, row["source"])
        _data_style(c_src, bg=bg)

        c_amt = ws.cell(r, 2, round(row["amount"], 2))
        _data_style(c_amt, bg=bg, color=_GREEN_H, bold=True, align="right")
        c_amt.number_format = '#,##0.00 "€"'

        c_lbl = ws.cell(r, 3, row["label"] or "")
        _data_style(c_lbl, bg=bg, color=_MUTED)

    # Ligne total
    if revenues:
        tr = len(revenues) + 3
        ws.row_dimensions[tr].height = 22
        ws.cell(tr, 1, "TOTAL").font = Font(bold=True, size=11, name="Calibri")
        ws.cell(tr, 1).fill = PatternFill("solid", fgColor=_GREEN_L)
        ws.cell(tr, 1).alignment = Alignment(horizontal="right")

        total_cell = ws.cell(tr, 2, f"=SUM(B3:B{tr-1})")
        _data_style(total_cell, bg=_GREEN_L, color=_GREEN_H, bold=True, align="right")
        total_cell.number_format = '#,##0.00 "€"'

        ws.cell(tr, 3).fill = PatternFill("solid", fgColor=_GREEN_L)

    ws.freeze_panes = "A3"
    _autofit(ws)


# ─────────────────────────────────────────────────────────────
#  TEMPLATE VIERGE
# ─────────────────────────────────────────────────────────────
def generate_template_excel(db, year: int, month: int, dest_path: str,
                            sheet_type: str = "both"):
    """
    Génère un fichier .xlsx vierge pré-formaté pour la saisie manuelle.

    sheet_type :
      "expenses" → feuille Dépenses seulement (avec dropdown catégories + 30 lignes vides)
      "revenues" → feuille Revenus seulement  (15 lignes vides)
      "both"     → les deux feuilles (compatibilité)

    Retourne le chemin du fichier créé.
    """
    from config import MONTHS_FR
    from openpyxl.worksheet.datavalidation import DataValidation

    month_label = f"{MONTHS_FR[month - 1]} {year}"
    cats = db.get_categories()
    cat_names = [c["name"] for c in cats]

    wb = Workbook()
    NB_DATA_ROWS = 30  # lignes vides pour la saisie

    # ─────────────────────────────────────────────────────────
    #  Feuille CATÉGORIES (masquée) — source des dropdowns
    #  Permet de dépasser la limite de 255 caractères de formula1
    # ─────────────────────────────────────────────────────────
    ws_cat = wb.active
    ws_cat.title = "_Catégories"
    for i, name in enumerate(cat_names, 1):
        ws_cat.cell(i, 1, name)
    ws_cat.sheet_state = "hidden"

    cat_range = f"_Catégories!$A$1:$A${max(len(cat_names), 1)}"

    # ─────────────────────────────────────────────────────────
    #  Feuille DÉPENSES
    # ─────────────────────────────────────────────────────────
    if sheet_type in ("expenses", "both"):
        ws_exp = wb.create_sheet("Dépenses")
        ws_exp.sheet_view.showGridLines = False
        ws_exp.row_dimensions[1].height = 36

        # Titre
        ws_exp.merge_cells("A1:D1")
        t = ws_exp["A1"]
        t.value = f"💸  Dépenses — {month_label}"
        t.font  = Font(bold=True, size=14, color=_RED_H, name="Calibri")
        t.fill  = PatternFill("solid", fgColor=_RED_L)
        t.alignment = Alignment(horizontal="center", vertical="center")

        # En-têtes
        headers = ["Catégorie ▼", "Enseigne", "Montant (€)", "Description"]
        for col, h in enumerate(headers, 1):
            c = ws_exp.cell(row=2, column=col, value=h)
            _hdr_style(c, bg=_RED_H)
        ws_exp.row_dimensions[2].height = 26

        # Instruction
        ws_exp.merge_cells("A3:D3")
        hi = ws_exp["A3"]
        hi.value = "ℹ  Sélectionnez la catégorie via le menu déroulant ▼  ·  Laissez Montant à 0 pour ignorer la ligne"
        hi.font  = Font(italic=True, size=10, color=_MUTED, name="Calibri")
        hi.fill  = PatternFill("solid", fgColor="FFFBEB")
        hi.alignment = Alignment(horizontal="center", vertical="center")
        ws_exp.row_dimensions[3].height = 16

        # Lignes de saisie
        for i in range(NB_DATA_ROWS):
            r = i + 4
            bg = "FFFFFF" if i % 2 == 0 else "FEF9FA"
            ws_exp.row_dimensions[r].height = 22
            c_cat = ws_exp.cell(r, 1, "")
            _data_style(c_cat, bg=bg)
            c_pay = ws_exp.cell(r, 2, "")
            _data_style(c_pay, bg=bg, color=_MUTED)
            c_amt = ws_exp.cell(r, 3, 0)
            _data_style(c_amt, bg=bg, color=_MUTED, bold=False, align="right")
            c_amt.number_format = '#,##0.00 "€"'
            c_lbl = ws_exp.cell(r, 4, "")
            _data_style(c_lbl, bg=bg, color=_MUTED)

        # Dropdown catégories (colonne A, lignes 4 à 4+NB_DATA_ROWS)
        dv = DataValidation(
            type="list",
            formula1=cat_range,
            allow_blank=True,
            showDropDown=False,   # False = flèche visible
            showErrorMessage=True,
            errorTitle="Catégorie invalide",
            error="Choisissez une catégorie dans la liste.",
        )
        ws_exp.add_data_validation(dv)
        dv.sqref = f"A4:A{3 + NB_DATA_ROWS}"

        ws_exp.freeze_panes = "A4"
        _autofit(ws_exp)
        ws_exp.column_dimensions["A"].width = 22
        ws_exp.column_dimensions["B"].width = 20
        ws_exp.column_dimensions["C"].width = 14
        ws_exp.column_dimensions["D"].width = 30

    # ─────────────────────────────────────────────────────────
    #  Feuille REVENUS
    # ─────────────────────────────────────────────────────────
    if sheet_type in ("revenues", "both"):
        ws_rev = wb.create_sheet("Revenus")
        ws_rev.sheet_view.showGridLines = False
        ws_rev.row_dimensions[1].height = 36

        ws_rev.merge_cells("A1:C1")
        t2 = ws_rev["A1"]
        t2.value = f"💶  Revenus — {month_label}"
        t2.font  = Font(bold=True, size=14, color=_GREEN_H, name="Calibri")
        t2.fill  = PatternFill("solid", fgColor=_GREEN_L)
        t2.alignment = Alignment(horizontal="center", vertical="center")

        headers2 = ["Source", "Montant (€)", "Description"]
        for col, h in enumerate(headers2, 1):
            c = ws_rev.cell(row=2, column=col, value=h)
            _hdr_style(c, bg=_GREEN_H)
        ws_rev.row_dimensions[2].height = 26

        ws_rev.merge_cells("A3:C3")
        hi2 = ws_rev["A3"]
        hi2.value = "ℹ  Saisissez vos sources de revenus (Salaire, Freelance, Loyer perçu…)  ·  Montant = 0 → ligne ignorée"
        hi2.font  = Font(italic=True, size=10, color=_MUTED, name="Calibri")
        hi2.fill  = PatternFill("solid", fgColor="FFFBEB")
        hi2.alignment = Alignment(horizontal="center", vertical="center")
        ws_rev.row_dimensions[3].height = 16

        for i in range(15):
            r = i + 4
            bg = "FFFFFF" if i % 2 == 0 else "F0FDF9"
            ws_rev.row_dimensions[r].height = 22
            c_src = ws_rev.cell(r, 1, "")
            _data_style(c_src, bg=bg)
            c_amt = ws_rev.cell(r, 2, 0)
            _data_style(c_amt, bg=bg, color=_MUTED, bold=False, align="right")
            c_amt.number_format = '#,##0.00 "€"'
            c_lbl = ws_rev.cell(r, 3, "")
            _data_style(c_lbl, bg=bg, color=_MUTED)

        ws_rev.freeze_panes = "A4"
        ws_rev.column_dimensions["A"].width = 24
        ws_rev.column_dimensions["B"].width = 14
        ws_rev.column_dimensions["C"].width = 30

    os.makedirs(os.path.dirname(os.path.abspath(dest_path)), exist_ok=True)
    wb.save(dest_path)
    log.info("Template Excel (%s) généré : %s", sheet_type, dest_path)
    return dest_path


# ─────────────────────────────────────────────────────────────
#  EXPORT
# ─────────────────────────────────────────────────────────────
def export_month_excel(db, year: int, month: int, dest_path: str):
    """
    Génère un fichier .xlsx avec les dépenses et revenus du mois.
    Retourne le chemin du fichier créé.
    """
    from config import MONTHS_FR
    month_label = f"{MONTHS_FR[month - 1]} {year}"

    expenses = db.get_expenses(year, month)
    revenues = db.get_revenues(year, month)

    wb = Workbook()
    ws_exp = wb.active
    _write_expenses_sheet(ws_exp, expenses, month_label)

    ws_rev = wb.create_sheet()
    _write_revenues_sheet(ws_rev, revenues, month_label)

    os.makedirs(os.path.dirname(os.path.abspath(dest_path)), exist_ok=True)
    wb.save(dest_path)
    log.info("Export Excel : %s (%d dépenses, %d revenus)",
             dest_path, len(expenses), len(revenues))
    return dest_path


# ─────────────────────────────────────────────────────────────
#  IMPORT
# ─────────────────────────────────────────────────────────────
def import_month_excel(db, year: int, month: int, src_path: str) -> dict:
    """
    Lit un fichier .xlsx exporté par Fintrack et importe les lignes
    dans la base de données.

    Retourne un dict :
    {
        "expenses_added": int,
        "revenues_added": int,
        "errors": list[str],
    }
    """
    result = {"expenses_added": 0, "revenues_added": 0, "errors": []}

    try:
        wb = load_workbook(src_path, data_only=True)
    except Exception as e:
        result["errors"].append(f"Impossible d'ouvrir le fichier : {e}")
        return result

    cats    = db.get_categories()
    cat_map = {c["name"].upper(): c["id"] for c in cats}

    # ── Feuille Dépenses ─────────────────────────────────────
    ws_exp = wb["Dépenses"] if "Dépenses" in wb.sheetnames else None
    if ws_exp:
        for row in ws_exp.iter_rows(min_row=3, values_only=True):
            cat_raw, payee, amount, label = (row + (None, None, None, None))[:4]
            # Ignorer les lignes vides, d'aide (ℹ️) ou d'en-tête répété
            if not cat_raw:
                continue
            if str(cat_raw).strip().startswith("ℹ"):
                continue
            if not amount:
                continue
            try:
                amount = float(str(amount).replace(",", ".").replace("€", "").strip())
            except ValueError:
                result["errors"].append(f"Montant invalide (dépense) : {amount}")
                continue
            if amount <= 0:
                continue

            cat_name = str(cat_raw).strip().upper()
            # Créer la catégorie si elle n'existe pas
            if cat_name not in cat_map:
                db.add_category(cat_name)
                cats   = db.get_categories()
                cat_map = {c["name"].upper(): c["id"] for c in cats}

            cat_id = cat_map.get(cat_name)
            if not cat_id:
                result["errors"].append(f"Catégorie introuvable : {cat_raw}")
                continue

            db.add_expense(
                year, month, cat_id,
                amount,
                str(label).strip() if label else "",
                str(payee).strip().upper() if payee else "",
            )
            result["expenses_added"] += 1
    else:
        result["errors"].append("Feuille 'Dépenses' introuvable dans le fichier.")

    # ── Feuille Revenus ──────────────────────────────────────
    ws_rev = wb["Revenus"] if "Revenus" in wb.sheetnames else None
    if ws_rev:
        for row in ws_rev.iter_rows(min_row=3, values_only=True):
            source, amount, label = (row + (None, None, None))[:3]
            if not source:
                continue
            if str(source).strip().startswith("ℹ"):
                continue
            if not amount:
                continue
            try:
                amount = float(str(amount).replace(",", ".").replace("€", "").strip())
            except ValueError:
                result["errors"].append(f"Montant invalide (revenu) : {amount}")
                continue
            if amount <= 0:
                continue

            db.add_revenue(
                year, month,
                str(source).strip(),
                amount,
                str(label).strip() if label else "",
            )
            result["revenues_added"] += 1
    else:
        result["errors"].append("Feuille 'Revenus' introuvable dans le fichier.")

    log.info("Import Excel : %d dépenses, %d revenus, %d erreurs",
             result["expenses_added"], result["revenues_added"], len(result["errors"]))
    return result
